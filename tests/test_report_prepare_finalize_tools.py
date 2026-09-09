from __future__ import annotations

import hashlib
import copy
import json
import re
import tempfile
import threading
import unittest
from pathlib import Path
from typing import Any
from unittest.mock import patch

import jsonschema
from docx import Document
from docx.oxml.ns import qn

from src.errors import HostStorageError, HostStorageIntegrityError
from src.report_finalize import OperationCancelled as FinalizeCancelled
from src.report_finalize import execute as execute_finalize
from src.report_prepare import OperationCancelled as PrepareCancelled
from src.report_prepare import execute as execute_prepare
from src.report_shared.workflow import (
    MANIFEST_SCHEMA_PATH,
    PACKAGE_SCHEMA_PATH,
    RESULTS_SCHEMA_PATH,
    SUMMARY_FIELDS,
)
from src.report_shared.context_builder import FACT_ROOTS
from src.report_shared.template_loader import (
    CHAPTER_RULES_PATH,
    STANDARDS_LIBRARY_PATH,
    TEMPLATE_PATH,
    load_template,
)
from tests.fakes import FakeContent, FakeHostClient


TOOL_ROOT = Path(__file__).resolve().parents[1]
EXISTING_DRAFT_TITLES_PATH = TOOL_ROOT / "tests" / "fixtures" / "existing_draft_section_titles.txt"
EXPECTED_WRITING_IDS = [
    "1.1.2", "1.1.3", "2.1.1", "2.1.2", "2.2", "2.4", "4.1.2",
    "4.1.3.1", "4.1.3.2", "4.1.3.3", "4.1.4", "4.2.1", "18.2",
]
EXPECTED_SYNTHESIS_IDS = ["1.2", "26.1.1", "26.1.2", "26.1.3", "26.2", "26.3", "26.4"]


def minimal_facts() -> dict[str, Any]:
    return {
        "meta": {"schema_version": "2.0"},
        "sources": [],
        "basic_info": {"construction_unit": "测试建设单位"},
        "unit": {
            "name": "测试装置",
            "annual_operating_hours": 8000,
            "design_case": {
                "products": [
                    {
                        "name": "产品A",
                        "stream_id": "P1",
                        "annual_capacity": {"value": 10000, "unit": "吨/年"},
                        "hourly_rate": {"value": 1250, "unit": "kg/h"},
                        "is_main_product": True,
                    }
                ]
            },
            "retrofit_case": {
                "products": [
                    {
                        "name": "产品A",
                        "stream_id": "P1",
                        "annual_capacity": {"value": 20000, "unit": "吨/年"},
                        "hourly_rate": {"value": 2500, "unit": "kg/h"},
                        "is_main_product": True,
                    }
                ]
            },
        },
        "component_catalog": [{"name": "产品A", "formula": "A"}],
        "process": {"design": {"streams": []}, "retrofit": {"streams": []}},
        "diagnosis": {"bottlenecks": ["测试瓶颈"]},
        "scheme_analysis": {
            "candidates": [{"name": "方案A", "retrofit_type": "add_equipment", "brief": "新增设备"}],
            "recommendation": {"name": "方案A"},
        },
        "adopted_scheme": {
            "selected_names": ["方案A"],
            "optimized": True,
            "schemes": [{"name": "方案A", "brief": "新增设备", "description": "新增设备。"}],
        },
        "equipment": {"object_catalog": []},
        "derived_facts": {},
    }


def empty_summary() -> dict[str, list[str]]:
    return {field: [] for field in SUMMARY_FIELDS}


def implementation_schedule() -> dict[str, Any]:
    return {
        "rows": [
            {
                "stage": stage,
                "main_content": "完成对应阶段工作。",
                "duration_range": "1～2个月",
                "prerequisites": "前序工作完成",
            }
            for stage in [
                "项目前期（各报告编制及审批）",
                "基础设计",
                "施工图设计",
                "设备采购",
                "土建施工",
                "安装工程",
                "试生产",
            ]
        ]
    }


def equipment_facts() -> dict[str, Any]:
    """Facts fixture exercising selected equipment actions and comparisons."""
    facts = minimal_facts()
    facts["unit"]["type"] = "联合生产装置"
    facts["scheme_analysis"] = {
        "candidates": [
            {
                "name": "方案A",
                "retrofit_type": "add_equipment",
                "brief": "新增设备",
                "evaluation": {
                    "technical_feasibility": {"level": "待核实"},
                    "implementation_complexity": {"level": "中"},
                    "operational_risk": {"level": "待核实"},
                    "energy_efficiency": {"comment": "能耗影响待核实", "display_name": "能效影响"},
                },
            }
        ],
        "user_candidates": [
            {
                "name": "自定义方案",
                "source_pool": "user_candidates",
                "retrofit_type": "equipment_service_change",
                "brief": "优化内件",
                "evaluation": {"maintainability": {"value": "中"}},
            }
        ],
        "recommendation": {"name": "方案A"},
    }
    facts["equipment"] = {
        "object_catalog": [
            {"id": "1", "name": "既有反应器", "type": "reactor", "tag": "R-101"},
            {"id": "2", "name": "既有精馏塔", "type": "separator"},
            {"id": "2_1", "name": "既有精馏塔_并联", "type": "separator", "is_new": True},
            {"id": "3", "name": "新反应器", "type": "reactor", "is_new": True},
            {"id": "4", "name": "并联反应器", "type": "reactor"},
        ],
        "reactor": {
            "selected_scheme": {
                "scheme_id": "combined_1",
                "scheme_index": 1,
                "reuse": [{"id": "3", "name": "新反应器", "scheme_type": "reuse", "reused_from": "1"}],
                "addition": [{"id": "1", "name": "既有反应器", "scheme_type": "increase_volume"}],
                "parallel": [{"id": "4", "name": "并联反应器", "scheme_type": "add_parallel"}],
            },
            "combined_schemes": [
                {"scheme_id": "combined_1", "scheme_index": 1, "conclusion": "方案1：既有反应器扩容，并联反应器新增"},
                {"scheme_id": "combined_2", "scheme_index": 2, "conclusion": "方案2：全部采用新增并联反应器"},
            ],
            "evaluations": [
                {
                    "id": "1", "tag": "R-101", "name": "既有反应器", "form": "固定床",
                    "retrofit": {"schemes": [{"scheme_type": "increase_volume", "description": "扩容方案：更换更大规格反应器", "diameter": 1800, "cylinder_height": 12000, "total_volume": 40.0, "catalyst_volume": 24.0, "quantity": 1}]},
                },
                {
                    "id": "3", "name": "新反应器", "form": "固定床",
                    "retrofit": {"schemes": [{"scheme_type": "reuse", "description": "利旧方案：复用既有反应器", "diameter": 1600, "cylinder_height": 9000, "catalyst_volume": 18.0, "quantity": 1}]},
                },
                {
                    "id": "4", "tag": "R-104", "name": "并联反应器", "form": "列管式",
                    "retrofit": {"schemes": [{"scheme_type": "add_parallel", "description": "并联方案：增设并联反应器", "diameter": 1000, "cylinder_height": 6000, "catalyst_volume": 9.0, "heat_exchange_area": 80, "quantity": 2}]},
                },
            ],
        },
        "tower": {
            "combined_schemes": {"optimal": [], "parallel": [[{"id": "2", "name": "既有精馏塔"}, {"id": "2_1", "name": "既有精馏塔_并联", "is_new": True}]], "series": [], "other": []},
            "evaluations": [{"id": "2", "name": "既有精馏塔", "retrofit": {"route": "parallel", "plan_detail": {"device_paras": {"struct_info": [[800, 300], [1000, 450]], "tray_num": 24, "tray_type": "F1"}}}}],
        },
    }
    return facts


class ReportPrepareFinalizeTests(unittest.TestCase):
    def setUp(self) -> None:
        self.host = FakeHostClient()
        self.content = FakeContent()
        self.host.save_file("inputs/engineering_facts.json", minimal_facts())

    def prepare(
        self,
        facts: dict[str, Any] | None = None,
        report_context: dict[str, Any] | None = None,
    ) -> tuple[dict[str, Any], dict[str, Any]]:
        if facts is not None:
            self.host.save_file("inputs/engineering_facts.json", facts)
        request: dict[str, Any] = {
            "engineering_facts_path": "inputs/engineering_facts.json",
            "report_context": (
                {"project_name": "测试项目"} if report_context is None else report_context
            ),
        }
        result = execute_prepare(
            request,
            content=self.content,
            host_client=self.host,
        )
        self.assertEqual(result["status"], "prepared", result)
        package_path = result["artifact"]["path"]
        package = self.host.get_file(package_path)
        self.assertIsInstance(package, dict)
        return result, package

    def context_for(self, package: dict[str, Any], section_id: str) -> dict[str, Any]:
        task = next(item for item in package["writing_tasks"] if item["section_id"] == section_id)
        payload = self.host.get_file(task["context_path"])
        self.assertIsInstance(payload, dict)
        return payload

    def finalize_with_results(
        self,
        package_result: dict[str, Any],
        package: dict[str, Any],
        results: dict[str, Any] | str,
        *,
        path: str = "inputs/work_results.json",
    ) -> dict[str, Any]:
        self.host.save_file(path, results)
        return execute_finalize(
            {
                "work_package_path": package_result["artifact"]["path"],
                "work_results_path": path,
            },
            content=self.content,
            host_client=self.host,
        )

    def valid_results(self, package: dict[str, Any]) -> dict[str, Any]:
        return {
            "schema_version": "1.0",
            "writing_results": [
                {
                    "section_id": task["section_id"],
                    "blocks": [{"type": "paragraph", "text": f"{task['section_id']} Agent 段落。"}],
                    **(
                        {"implementation_schedule": implementation_schedule()}
                        if task["section_id"] == "18.2"
                        else {}
                    ),
                    "section_summary": empty_summary(),
                }
                for task in package["writing_tasks"]
            ],
            "synthesis_results": [
                {
                    "section_id": task["section_id"],
                    "blocks": [{"type": "paragraph", "text": f"{task['section_id']} 综合段落。"}],
                    "section_summary": empty_summary(),
                }
                for task in package["synthesis_tasks"]
            ],
        }

    def test_prepare_commits_contexts_before_work_package_with_path_schema(self) -> None:
        result, package = self.prepare()

        jsonschema.validate(package, json.loads(PACKAGE_SCHEMA_PATH.read_text(encoding="utf-8")))
        self.assertIn("/work_package.json", result["artifact"]["path"])
        self.assertEqual(package["engineering_facts"]["path"], "inputs/engineering_facts.json")
        context_paths = [task["context_path"] for task in package["writing_tasks"]]
        self.assertTrue(context_paths)
        for context_path in context_paths:
            self.assertIn(context_path, self.host.files)
        self.assertIn(result["artifact"]["path"], self.host.files)

    def test_prepare_context_save_failure_does_not_save_work_package(self) -> None:
        self.host.fail_on_save_prefix = "runs/report_prepare"

        with self.assertRaises(HostStorageError):
            execute_prepare(
                {"engineering_facts_path": "inputs/engineering_facts.json"},
                content=self.content,
                host_client=self.host,
            )

        self.assertFalse([path for path in self.host.files if path.endswith("/work_package.json")])

    def test_finalize_without_results_keeps_fallback_draft_and_manifest(self) -> None:
        prepare_result, _ = self.prepare()

        result = execute_finalize(
            {"work_package_path": prepare_result["artifact"]["path"]},
            content=self.content,
            host_client=self.host,
        )

        self.assertEqual(result["status"], "completed", result)
        self.assertGreater(result["summary"]["fallback_section_count"], 0)
        self.assertTrue(result["manifest_path"].endswith("/report_manifest.json"))
        self.assertIn(result["manifest_path"], self.host.files)
        self.assertIn("WORK_RESULTS_NOT_PROVIDED", {d["code"] for d in result["diagnostics"]})
        manifest = self.host.get_file(result["manifest_path"])
        self.assertIsInstance(manifest, dict)
        jsonschema.validate(manifest, json.loads(MANIFEST_SCHEMA_PATH.read_text(encoding="utf-8")))
        self.assertEqual(manifest["markdown"]["path"], result["artifacts"]["markdown_path"])
        self.assertEqual(manifest["docx"]["path"], result["artifacts"]["docx_path"])
        for field in ("markdown", "docx"):
            payload = self.host.get_file(manifest[field]["path"])
            self.assertIsInstance(payload, bytes)
            self.assertEqual(manifest[field]["size_bytes"], len(payload))
            self.assertEqual(manifest[field]["sha256"], hashlib.sha256(payload).hexdigest())

    def test_finalize_with_results_exports_agent_content_and_valid_manifest(self) -> None:
        prepare_result, package = self.prepare()
        self.host.save_file("inputs/work_results.json", self.valid_results(package))

        result = execute_finalize(
            {
                "work_package_path": prepare_result["artifact"]["path"],
                "work_results_path": "inputs/work_results.json",
            },
            content=self.content,
            host_client=self.host,
        )

        self.assertEqual(result["status"], "completed", result)
        self.assertEqual(result["summary"]["fallback_section_count"], 0)
        self.assertIn("Agent 段落", result["markdown_content"])
        jsonschema.validate(
            self.host.get_file(result["manifest_path"]),
            json.loads(MANIFEST_SCHEMA_PATH.read_text(encoding="utf-8")),
        )
        jsonschema.validate(
            self.host.get_file("inputs/work_results.json"),
            json.loads(RESULTS_SCHEMA_PATH.read_text(encoding="utf-8")),
        )

    def test_finalize_rejects_inline_work_results(self) -> None:
        prepare_result, _ = self.prepare()

        result = execute_finalize(
            {
                "work_package_path": prepare_result["artifact"]["path"],
                "work_results": {"schema_version": "1.0"},
            },
            content=self.content,
            host_client=self.host,
        )

        self.assertEqual(result["status"], "failed")
        self.assertIsNone(result["manifest_path"])

    def test_finalize_storage_get_failure_raises_tool_boundary_error(self) -> None:
        prepare_result, _ = self.prepare()
        self.host.fail_on_get_prefix = prepare_result["artifact"]["path"]

        with self.assertRaises(HostStorageError):
            execute_finalize(
                {"work_package_path": prepare_result["artifact"]["path"]},
                content=self.content,
                host_client=self.host,
            )

    def test_finalize_cover_storage_failure_is_not_silently_downgraded(self) -> None:
        prepare_result, _ = self.prepare()
        self.host.fail_on_get_prefix = "inputs/engineering_facts.json"

        with self.assertRaises(HostStorageError):
            execute_finalize(
                {"work_package_path": prepare_result["artifact"]["path"]},
                content=self.content,
                host_client=self.host,
            )

    def test_prepare_cancelled_before_start_raises_and_commits_nothing(self) -> None:
        cancel_event = threading.Event()
        cancel_event.set()

        with self.assertRaises(PrepareCancelled):
            execute_prepare(
                {"engineering_facts_path": "inputs/engineering_facts.json"},
                content=self.content,
                host_client=self.host,
                cancel_event=cancel_event,
            )
        self.assertFalse([path for path in self.host.files if path.startswith("runs/report_prepare")])

    def test_prepare_cancelled_at_commit_marker_does_not_save_work_package(self) -> None:
        cancel_event = threading.Event()

        class CancellingContent(FakeContent):
            def report_progress(self, progress, total=None, message=None):
                super().report_progress(progress, total, message)
                if progress == 90:
                    cancel_event.set()

        with self.assertRaises(PrepareCancelled):
            execute_prepare(
                {"engineering_facts_path": "inputs/engineering_facts.json"},
                content=CancellingContent(),
                host_client=self.host,
                cancel_event=cancel_event,
            )

        self.assertFalse([path for path in self.host.files if path.endswith("/work_package.json")])

    def test_finalize_cancelled_before_manifest_commits_no_manifest(self) -> None:
        prepare_result, _ = self.prepare()
        cancel_event = threading.Event()
        cancel_event.set()

        with self.assertRaises(FinalizeCancelled):
            execute_finalize(
                {"work_package_path": prepare_result["artifact"]["path"]},
                content=self.content,
                host_client=self.host,
                cancel_event=cancel_event,
            )
        self.assertFalse([path for path in self.host.files if path.endswith("/report_manifest.json")])

    def test_finalize_cancelled_at_manifest_checkpoint_commits_no_manifest(self) -> None:
        prepare_result, _ = self.prepare()
        cancel_event = threading.Event()

        class CancellingContent(FakeContent):
            def report_progress(self, progress, total=None, message=None):
                super().report_progress(progress, total, message)
                if progress == 95:
                    cancel_event.set()

        with self.assertRaises(FinalizeCancelled):
            execute_finalize(
                {"work_package_path": prepare_result["artifact"]["path"]},
                content=CancellingContent(),
                host_client=self.host,
                cancel_event=cancel_event,
            )

        self.assertFalse([path for path in self.host.files if path.endswith("/report_manifest.json")])

    def test_finalize_readback_mismatch_is_internal_and_commits_no_manifest(self) -> None:
        prepare_result, _ = self.prepare()

        class CorruptingHostClient(FakeHostClient):
            def get_file(self, path, *, kind="auto"):
                payload = super().get_file(path, kind=kind)
                if path.endswith(".md") and isinstance(payload, bytes):
                    return payload + b"corrupted"
                return payload

        host = CorruptingHostClient()
        host.files = self.host.files.copy()

        with self.assertRaises(HostStorageIntegrityError):
            execute_finalize(
                {"work_package_path": prepare_result["artifact"]["path"]},
                content=FakeContent(),
                host_client=host,
            )

        self.assertFalse([path for path in host.files if path.endswith("/report_manifest.json")])

    def test_prepare_generates_complete_schema_valid_task_graph(self) -> None:
        result, package = self.prepare(report_context={})
        self.assertEqual(package["status"], "prepared")
        self.assertEqual(package["project_name"], "测试建设单位测试装置技术改造项目")
        self.assertEqual([section["order"] for section in package["sections"]], list(range(1, 67)))
        self.assertEqual([task["section_id"] for task in package["writing_tasks"]], EXPECTED_WRITING_IDS)
        self.assertEqual([task["section_id"] for task in package["synthesis_tasks"]], EXPECTED_SYNTHESIS_IDS)
        self.assertEqual(result["summary"]["writing_task_count"], 13)
        self.assertEqual(result["summary"]["synthesis_task_count"], 7)
        self.assertLessEqual(result["summary"]["research_task_count"], 4)
        jsonschema.validate(package, json.loads(PACKAGE_SCHEMA_PATH.read_text(encoding="utf-8")))

    def test_prepare_is_deterministic_for_identical_inputs(self) -> None:
        with patch("src.report_prepare.core.new_logical_prefix", return_value="runs/report_prepare/fixed"):
            _, first = self.prepare(report_context={})
            first_files = copy.deepcopy(self.host.files)
            _, second = self.prepare(report_context={})
        self.assertEqual(first, second)
        self.assertEqual(first["created_at"], "1970-01-01T00:00:00+00:00")
        self.assertEqual(first_files["runs/report_prepare/fixed/work_package.json"], second)

    def test_research_scope_is_narrow_and_optional_by_section(self) -> None:
        _, package = self.prepare()
        groups = {task["research_group"]: task["target_section_ids"] for task in package["research_tasks"]}
        self.assertLessEqual(len(groups), 4)
        self.assertEqual(groups["technology_research"], ["4.1.2"])
        writing = {task["section_id"]: task for task in package["writing_tasks"]}
        self.assertIsNone(writing["1.1.3"]["research_task_id"])
        self.assertIsNone(writing["4.1.3.2"]["research_task_id"])

    def test_synthesis_sources_are_single_round_and_closed(self) -> None:
        _, package = self.prepare()
        writing_ids = {task["section_id"] for task in package["writing_tasks"]}
        synthesis_ids = {task["section_id"] for task in package["synthesis_tasks"]}
        deterministic_ids = set(package["deterministic_summaries"])
        referenced = set()
        for task in package["synthesis_tasks"]:
            for source_id in task["context"]["source_section_ids"]:
                self.assertNotIn(source_id, synthesis_ids)
                self.assertIn(source_id, writing_ids | deterministic_ids)
                if source_id not in writing_ids:
                    referenced.add(source_id)
        self.assertEqual(referenced, deterministic_ids)

    def test_project_name_uses_four_level_priority(self) -> None:
        _, package = self.prepare(report_context={"project_name": "显式项目名"})
        self.assertEqual(package["project_name"], "显式项目名")
        _, package = self.prepare(minimal_facts(), report_context={})
        self.assertEqual(package["project_name"], "测试建设单位测试装置技术改造项目")
        facts = minimal_facts()
        facts["basic_info"]["construction_unit"] = ""
        _, package = self.prepare(facts, report_context={})
        self.assertEqual(package["project_name"], "测试装置技术改造项目")
        facts["basic_info"] = {}
        facts["unit"] = {}
        _, package = self.prepare(facts, report_context={})
        self.assertEqual(package["project_name"], "工业装置技术改造项目")

    def test_missing_top_level_fact_root_fails_without_commit(self) -> None:
        facts = minimal_facts()
        facts.pop("equipment")
        self.host.save_file("inputs/engineering_facts.json", facts)
        result = execute_prepare({"engineering_facts_path": "inputs/engineering_facts.json"}, content=self.content, host_client=self.host)
        self.assertEqual(result["status"], "failed")
        self.assertIn("missing required roots", result["diagnostics"][0]["message"])
        self.assertFalse([path for path in self.host.files if path.endswith("/work_package.json")])

    def test_root_internal_fields_missing_still_prepares(self) -> None:
        facts = {"meta": {"schema_version": "2.0"}, "sources": [], "basic_info": {}, "unit": {}, "component_catalog": [], "process": {}, "diagnosis": {}, "scheme_analysis": {}, "adopted_scheme": {}, "equipment": {}, "derived_facts": {}}
        _, package = self.prepare(facts, report_context={})
        self.assertEqual(package["engineering_facts"]["root_keys"], FACT_ROOTS)
        self.assertEqual(package["project_name"], "工业装置技术改造项目")

    def test_invalid_top_level_fact_contract_fails(self) -> None:
        cases = [
            (lambda facts: facts.update({"project": {}}), "unsupported roots"),
            (lambda facts: facts["meta"].update({"schema_version": "1.0"}), "schema_version"),
            (lambda facts: facts.update({"sources": {}}), "invalid type"),
        ]
        for mutate, expected in cases:
            with self.subTest(expected=expected):
                facts = minimal_facts()
                mutate(facts)
                self.host.save_file("inputs/engineering_facts.json", facts)
                result = execute_prepare({"engineering_facts_path": "inputs/engineering_facts.json"}, content=self.content, host_client=self.host)
                self.assertEqual(result["status"], "failed")
                self.assertIn(expected, result["diagnostics"][0]["message"])

    def test_fact_slices_are_isolated_by_section_contract(self) -> None:
        _, package = self.prepare(equipment_facts())
        self.assertEqual(set(self.context_for(package, "1.1.2")["fact_slice"]), {"basic_info"})
        slice_4132 = self.context_for(package, "4.1.3.2")["fact_slice"]
        self.assertEqual(set(slice_4132), {"scheme_analysis"})
        self.assertEqual(set(slice_4132["scheme_analysis"]), {"candidates", "user_candidates"})
        slice_414 = self.context_for(package, "4.1.4")["fact_slice"]
        self.assertEqual(set(slice_414), {"unit", "process", "diagnosis", "adopted_scheme", "equipment"})
        self.assertEqual(set(slice_414["equipment"]), {"reactor", "tower"})
        self.assertEqual(set(self.context_for(package, "18.2")["fact_slice"]["equipment"]), {"object_catalog"})

    def test_equipment_tables_use_structured_selected_actions(self) -> None:
        _, package = self.prepare(equipment_facts())
        section = next(item for item in package["sections"] if item["id"] == "4.2.4")
        tables = {block["caption"]: block for block in section["deterministic_blocks"] if block["type"] == "table"}
        self.assertEqual(set(tables), {"新增工艺设备汇总表", "利旧工艺设备汇总表", "改造工艺设备汇总表"})
        self.assertEqual([len(tables[name]["rows"]) for name in tables], [2, 1, 1])
        serialized = json.dumps(section, ensure_ascii=False)
        self.assertIn("既有精馏塔（并联新增）", serialized)
        self.assertIn("扩容方案：更换更大规格反应器", serialized)

    def test_scheme_comparison_uses_dynamic_evaluation_union(self) -> None:
        _, package = self.prepare(equipment_facts())
        section = next(item for item in package["sections"] if item["id"] == "4.1.3.2")
        table = next(block for block in section["deterministic_blocks"] if block.get("caption") == "候选工艺技术方案比选表")
        self.assertEqual(table["headers"][-2:], ["能效影响", "Maintainability"])
        self.assertEqual(table["rows"][0][-2:], ["能耗影响待核实", "待补充"])
        self.assertEqual(table["rows"][1][-1], "中")

    def test_key_equipment_comparison_marks_selected_and_alternative(self) -> None:
        _, package = self.prepare(equipment_facts())
        section = next(item for item in package["sections"] if item["id"] == "4.1.3.2")
        table = next(block for block in section["deterministic_blocks"] if block.get("caption") == "关键设备方案比较表")
        self.assertEqual(table["table_id"], "表4.2")
        self.assertEqual([row[-1] for row in table["rows"]], ["选定", "备选"])

    def test_bad_key_equipment_selection_warns_and_skips_table(self) -> None:
        facts = equipment_facts()
        facts["equipment"]["reactor"]["selected_scheme"].update({"scheme_id": "missing", "scheme_index": 99})
        result, package = self.prepare(facts)
        self.assertIn("KEY_EQUIPMENT_COMPARISON_SKIPPED", {item["code"] for item in result["diagnostics"]})
        section = next(item for item in package["sections"] if item["id"] == "4.1.3.2")
        self.assertNotIn("关键设备方案比较表", {block.get("caption") for block in section["deterministic_blocks"]})

    def test_synthesis_tasks_do_not_embed_full_section_text(self) -> None:
        _, package = self.prepare()
        task = next(item for item in package["synthesis_tasks"] if item["section_id"] == "26.2")
        self.assertEqual(set(task["context"]), {"source_section_ids", "summary_fields"})
        serialized = json.dumps(task, ensure_ascii=False)
        for forbidden in ("fact_slice", "deterministic_blocks", "full_text"):
            self.assertNotIn(forbidden, serialized)

    def test_prepare_rejects_undeclared_request_and_context_fields(self) -> None:
        for request, expected in [
            ({"engineering_facts_path": "inputs/engineering_facts.json", "project_type": "x"}, "unsupported fields"),
            ({"engineering_facts_path": "inputs/engineering_facts.json", "report_context": {"project_type": "x"}}, "report_context has unsupported fields"),
        ]:
            with self.subTest(expected=expected):
                result = execute_prepare(request, content=self.content, host_client=self.host)
                self.assertEqual(result["status"], "failed")
                self.assertIn(expected, result["diagnostics"][0]["message"])

    def test_finalize_rejects_undeclared_fields_and_blank_results_path(self) -> None:
        prepared, _ = self.prepare()
        for request, expected in [
            ({"work_package_path": prepared["artifact"]["path"], "repair": True}, "unsupported fields"),
            ({"work_package_path": prepared["artifact"]["path"], "work_results_path": "  "}, "work_results_path"),
        ]:
            with self.subTest(expected=expected):
                result = execute_finalize(request, content=self.content, host_client=self.host)
                self.assertEqual(result["status"], "failed")
                self.assertIn(expected, result["diagnostics"][0]["message"])

    def test_finalize_exports_markdown_and_docx_schedule(self) -> None:
        prepared, package = self.prepare(report_context={})
        result = self.finalize_with_results(prepared, package, self.valid_results(package))
        self.assertEqual(result["status"], "completed")
        markdown = self.host.get_file(result["artifacts"]["markdown_path"], kind="text")
        self.assertIn("测试建设单位测试装置技术改造项目", markdown)
        self.assertIn("**表18.1 项目实施进度计划表**", markdown)
        docx_bytes = self.host.get_file(result["artifacts"]["docx_path"], kind="bytes")
        from io import BytesIO
        doc = Document(BytesIO(docx_bytes))
        schedule_tables = [table for table in doc.tables if [cell.text for cell in table.rows[0].cells] == ["阶段", "主要内容", "预计时长", "前置条件"]]
        self.assertEqual(len(schedule_tables), 1)
        self.assertTrue(all(row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None for row in schedule_tables[0].rows))

    def test_fallback_report_contains_every_agent_section_fallback(self) -> None:
        prepared, package = self.prepare()
        result = execute_finalize({"work_package_path": prepared["artifact"]["path"]}, content=self.content, host_client=self.host)
        expected = len(package["writing_tasks"]) + len(package["synthesis_tasks"])
        self.assertEqual(result["summary"]["fallback_section_count"], expected)
        self.assertEqual([item["code"] for item in result["diagnostics"]].count("SECTION_FALLBACK"), expected)
        for section in package["sections"]:
            if section["agent_slot"] is not None:
                self.assertIn(section["fallback"], result["markdown_content"])

    def test_missing_and_invalid_dynamic_sections_use_fallback(self) -> None:
        prepared, package = self.prepare()
        results = self.valid_results(package)
        results["writing_results"] = [item for item in results["writing_results"] if item["section_id"] != "1.1.2"]
        invalid = next(item for item in results["writing_results"] if item["section_id"] == "1.1.3")
        invalid["blocks"] = [{"type": "table", "text": "非法块"}]
        result = self.finalize_with_results(prepared, package, results)
        self.assertEqual(result["summary"]["fallback_section_count"], 2)
        for section_id in ("1.1.2", "1.1.3"):
            fallback = next(item["fallback"] for item in package["sections"] if item["id"] == section_id)
            self.assertIn(fallback, result["markdown_content"])

    def test_empty_dynamic_blocks_use_fallback(self) -> None:
        prepared, package = self.prepare()
        results = self.valid_results(package)
        next(item for item in results["writing_results"] if item["section_id"] == "1.1.2")["blocks"] = []
        result = self.finalize_with_results(prepared, package, results)
        self.assertEqual(result["summary"]["fallback_section_count"], 1)

    def test_invalid_implementation_schedule_variants_fallback(self) -> None:
        cases = {
            "missing": lambda value: None,
            "missing_row": lambda value: {**value, "rows": value["rows"][:-1]},
            "duplicate": lambda value: {**value, "rows": [value["rows"][0], copy.deepcopy(value["rows"][0]), *value["rows"][2:]]},
            "wrong_order": lambda value: {**value, "rows": [value["rows"][1], value["rows"][0], *value["rows"][2:]]},
            "extra": lambda value: {**value, "rows": [*value["rows"], {**value["rows"][-1], "stage": "额外阶段"}]},
            "duration": lambda value: {**value, "rows": [{**value["rows"][0], "duration_range": "2个月"}, *value["rows"][1:]]},
        }
        for index, (name, mutate) in enumerate(cases.items()):
            with self.subTest(name=name):
                prepared, package = self.prepare()
                results = self.valid_results(package)
                target = next(item for item in results["writing_results"] if item["section_id"] == "18.2")
                changed = mutate(copy.deepcopy(target["implementation_schedule"]))
                if changed is None:
                    target.pop("implementation_schedule")
                else:
                    target["implementation_schedule"] = changed
                result = self.finalize_with_results(prepared, package, results, path=f"inputs/results_{index}.json")
                self.assertEqual(result["summary"]["fallback_section_count"], 1)
                self.assertNotIn("**表18.1 项目实施进度计划表**", result["markdown_content"])

    def test_implementation_schedule_is_rejected_outside_18_2(self) -> None:
        prepared, package = self.prepare()
        results = self.valid_results(package)
        next(item for item in results["writing_results"] if item["section_id"] == "1.1.2")["implementation_schedule"] = implementation_schedule()
        result = self.finalize_with_results(prepared, package, results)
        self.assertEqual(result["status"], "failed")
        self.assertIn("implementation_schedule", result["diagnostics"][0]["message"])

    def test_raw_optimized_marker_is_replaced_by_fallback(self) -> None:
        prepared, package = self.prepare()
        results = self.valid_results(package)
        target = next(item for item in results["writing_results"] if item["section_id"] == "4.1.3.3")
        target["blocks"] = [{"type": "paragraph", "text": "最终采用方案B，optimized=true。"}]
        result = self.finalize_with_results(prepared, package, results)
        self.assertEqual(result["summary"]["fallback_section_count"], 1)
        self.assertNotIn("optimized=true", result["markdown_content"])

    def test_missing_result_collection_fails(self) -> None:
        prepared, package = self.prepare()
        results = self.valid_results(package)
        del results["synthesis_results"]
        result = self.finalize_with_results(prepared, package, results)
        self.assertEqual(result["status"], "failed")
        self.assertIn("missing required fields", result["diagnostics"][0]["message"])

    def test_duplicate_or_unknown_result_section_id_fails(self) -> None:
        prepared, package = self.prepare()
        duplicate = self.valid_results(package)
        duplicate["writing_results"].append(copy.deepcopy(duplicate["writing_results"][0]))
        result = self.finalize_with_results(prepared, package, duplicate, path="inputs/duplicate.json")
        self.assertEqual(result["status"], "failed")
        self.assertIn("duplicate section_id", result["diagnostics"][0]["message"])
        unknown = self.valid_results(package)
        unknown["writing_results"][0]["section_id"] = "unknown"
        result = self.finalize_with_results(prepared, package, unknown, path="inputs/unknown.json")
        self.assertEqual(result["status"], "failed")
        self.assertIn("unknown section_id", result["diagnostics"][0]["message"])

    def test_agent_cannot_control_titles_or_table_numbers(self) -> None:
        prepared, package = self.prepare()
        results = self.valid_results(package)
        for item in results["writing_results"]:
            item["title"] = "Agent 伪标题"
            item["table_id"] = "Agent 伪表号"
        result = self.finalize_with_results(prepared, package, results)
        self.assertIn("### 1.1.2 主办单位基本情况", result["markdown_content"])
        self.assertNotIn("Agent 伪标题", result["markdown_content"])
        self.assertNotIn("Agent 伪表号", result["markdown_content"])

    def test_paragraph_bullet_numbered_blocks_export_to_both_formats(self) -> None:
        prepared, package = self.prepare()
        results = self.valid_results(package)
        target = next(item for item in results["writing_results"] if item["section_id"] == "1.1.2")
        target["blocks"] = [
            {"type": "paragraph", "text": "段落块导出检查。"},
            {"type": "bullet_list", "items": ["无序项A", "无序项B"]},
            {"type": "numbered_list", "items": ["有序项A", "有序项B"]},
        ]
        result = self.finalize_with_results(prepared, package, results)
        self.assertIn("- 无序项A", result["markdown_content"])
        self.assertIn("1. 有序项A", result["markdown_content"])
        from io import BytesIO
        doc = Document(BytesIO(self.host.get_file(result["artifacts"]["docx_path"], kind="bytes")))
        text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
        for expected in ("段落块导出检查。", "无序项A", "有序项A"):
            self.assertIn(expected, text)

    def test_bad_json_and_non_object_results_fail_business_envelope(self) -> None:
        prepared, package = self.prepare()
        for index, payload in enumerate(("{bad json", '["not", "object"]')):
            result = self.finalize_with_results(prepared, package, payload, path=f"inputs/bad_{index}.json")
            self.assertEqual(result["status"], "failed")
            self.assertIsNone(result["manifest_path"])

    def test_work_results_schema_accepts_metadata_but_fixes_block_shape(self) -> None:
        _, package = self.prepare()
        results = self.valid_results(package)
        for item in results["writing_results"]:
            item["title"] = "metadata"
            item["table_id"] = "metadata"
        jsonschema.validate(results, json.loads(RESULTS_SCHEMA_PATH.read_text(encoding="utf-8")))

    def test_template_covers_existing_titles_and_bounded_fallbacks(self) -> None:
        template = json.loads(TEMPLATE_PATH.read_text(encoding="utf-8"))
        if EXISTING_DRAFT_TITLES_PATH.is_file():
            titles = {section["title"] for section in template["sections"]}
            expected = [line.strip() for line in EXISTING_DRAFT_TITLES_PATH.read_text(encoding="utf-8").splitlines() if line.strip()]
            self.assertFalse([title for title in expected if title not in titles])
        self.assertFalse([section["id"] for section in template["sections"] if "order" in section or "contract" in section])
        forbidden = re.compile(r"满足规范|依托现有系统|依托现有罐区|依托既有|依托现有管网")
        self.assertFalse([section["id"] for section in template["sections"] if forbidden.search(section.get("fallback", ""))])

    def test_finalize_cancelled_at_progress_80_writes_no_outputs(self) -> None:
        prepared, _ = self.prepare()
        cancel_event = threading.Event()
        class CancelAt80(FakeContent):
            def report_progress(self, progress, total=None, message=None):
                super().report_progress(progress, total, message)
                if progress == 80:
                    cancel_event.set()
        with self.assertRaises(FinalizeCancelled):
            execute_finalize({"work_package_path": prepared["artifact"]["path"]}, content=CancelAt80(), host_client=self.host, cancel_event=cancel_event)
        self.assertFalse([path for path in self.host.files if path.endswith((".md", ".docx", "/report_manifest.json"))])

    def test_finalize_cancel_after_first_file_leaves_only_first_orphan(self) -> None:
        prepared, _ = self.prepare()
        cancel_event = threading.Event()
        class CancelAfterFirstSave(FakeHostClient):
            def save_file(self, path, data, *, kind="auto"):
                super().save_file(path, data, kind=kind)
                if path.endswith(".md"):
                    cancel_event.set()
        host = CancelAfterFirstSave()
        host.files = self.host.files.copy()
        with self.assertRaises(FinalizeCancelled):
            execute_finalize({"work_package_path": prepared["artifact"]["path"]}, content=FakeContent(), host_client=host, cancel_event=cancel_event)
        self.assertEqual(len([path for path in host.files if path.endswith(".md")]), 1)
        self.assertFalse([path for path in host.files if path.endswith((".docx", "/report_manifest.json"))])

    def test_utilities_and_energy_tables_are_dynamic(self) -> None:
        facts = minimal_facts()
        facts["derived_facts"] = {
            "utility_consumption_summary": {"status": "identified", "annual_operating_hours": 8000, "items": [{"status": "annualized", "condition": "retrofit", "equipment_id": "R-101", "equipment_name": "新增反应器", "medium": "electricity", "medium_code": "electricity", "medium_name": "电力", "annual_quantity": {"value": 80000, "unit": "kWh/a"}, "formula": "power_kw * annual_operating_hours"}]},
            "energy_conversion": {"status": "calculated", "conversion_type": "standard_coal", "standard": "GB/T 2589-2020", "total_unit": "t标准煤（tce）", "items": [{"status": "calculated", "medium": "electricity", "medium_code": "electricity", "medium_name": "电力", "condition": "retrofit", "equipment_id": "R-101", "equipment_name": "新增反应器", "quantity": 80000, "unit": "kWh", "rule_id": "SC-ELEC-001", "coefficient": 0.1229, "coefficient_unit": "kgce/kWh", "standard_coal_tce": 9.832}], "total_standard_coal_tce": 9.832},
        }
        _, package = self.prepare(facts)
        utilities = next(item for item in package["sections"] if item["id"] == "5.4")
        energy = next(item for item in package["sections"] if item["id"] == "10.5")
        self.assertIn("新增反应器", json.dumps(utilities, ensure_ascii=False))
        energy_text = json.dumps(energy, ensure_ascii=False)
        self.assertIn("9.83", energy_text)
        self.assertIn("SC-ELEC-001", energy_text)

    def test_report_tables_keep_missing_conversion_reasons(self) -> None:
        facts = minimal_facts()
        facts["derived_facts"] = {
            "utility_consumption_summary": {"status": "identified", "items": [{"status": "annualized", "condition": "retrofit", "equipment_name": "冷却设备", "medium": "CW", "medium_code": "cooling_water", "medium_name": "循环冷却水", "annual_quantity": {"value": 1200, "unit": "t/a"}, "formula": "mass_flow"}]},
            "energy_conversion": {"status": "blocked", "conversion_type": "standard_coal", "items": [{"status": "blocked", "reason": "conversion_factor_missing", "medium": "CW", "medium_code": "cooling_water", "medium_name": "循环冷却水", "quantity": 1200, "unit": "t"}]},
        }
        _, package = self.prepare(facts)
        serialized = json.dumps(package["sections"], ensure_ascii=False)
        self.assertIn("循环冷却水", serialized)
        self.assertIn("缺少折标系数", serialized)
        self.assertNotIn("conversion_factor_missing", serialized)

    def test_blocked_energy_table_does_not_show_zero_total(self) -> None:
        facts = minimal_facts()
        facts["derived_facts"]["energy_conversion"] = {"status": "blocked", "conversion_type": "standard_coal", "items": [{"status": "blocked", "reason": "conversion_factor_missing", "medium": "CW", "medium_code": "cooling_water", "medium_name": "循环冷却水", "quantity": 1200, "unit": "t"}]}
        _, package = self.prepare(facts)
        table = next(item for item in package["sections"] if item["id"] == "10.5")["deterministic_blocks"][0]
        self.assertEqual(len(table["rows"]), 1)
        self.assertNotIn("合计", json.dumps(table, ensure_ascii=False))
        self.assertNotIn("0.00", json.dumps(table, ensure_ascii=False))

    def test_energy_table_separates_condition_subtotals(self) -> None:
        facts = minimal_facts()
        facts["derived_facts"]["energy_conversion"] = {
            "status": "calculated", "conversion_type": "standard_coal",
            "items": [
                {"status": "calculated", "medium": "electricity", "medium_code": "electricity", "medium_name": "电力", "condition": "design", "quantity": 80000, "unit": "kWh", "rule_id": "SC-ELEC-001", "coefficient": 0.1229, "coefficient_unit": "kgce/kWh", "standard_coal_tce": 9.832},
                {"status": "calculated", "medium": "electricity", "medium_code": "electricity", "medium_name": "电力", "condition": "retrofit", "quantity": 160000, "unit": "kWh", "rule_id": "SC-ELEC-001", "coefficient": 0.1229, "coefficient_unit": "kgce/kWh", "standard_coal_tce": 19.664},
            ],
            "totals_by_condition": [{"condition": "design", "item_count": 1, "total_standard_coal_tce": 9.832}, {"condition": "retrofit", "item_count": 1, "total_standard_coal_tce": 19.664}],
        }
        _, package = self.prepare(facts)
        table = next(item for item in package["sections"] if item["id"] == "10.5")["deterministic_blocks"][0]
        self.assertEqual([row for row in table["rows"] if row[0] == "小计"], [["小计", "改造前", "", "", "", "9.83", "已折算"], ["小计", "改造后", "", "", "", "19.66", "已折算"]])
        self.assertNotIn("29.50", json.dumps(table, ensure_ascii=False))

    def test_no_utility_data_uses_single_gap_row(self) -> None:
        _, package = self.prepare(minimal_facts())
        sections = [next(item for item in package["sections"] if item["id"] == section_id) for section_id in ("5.4", "10.5")]
        serialized = json.dumps(sections, ensure_ascii=False)
        self.assertIn("未识别到可汇总公用工程/待补充", serialized)
        self.assertNotIn("蒸汽", serialized)
        self.assertNotIn("循环冷却水", serialized)

    def test_bad_engineering_facts_json_is_business_failure(self) -> None:
        self.host.save_file("inputs/engineering_facts.json", "{bad json")
        result = execute_prepare({"engineering_facts_path": "inputs/engineering_facts.json"}, content=self.content, host_client=self.host)
        self.assertEqual(result["status"], "failed")
        self.assertIn("valid JSON object", result["diagnostics"][0]["message"])

    def test_corrupt_shared_template_is_internal_exception(self) -> None:
        with patch("src.report_shared.workflow.load_template", side_effect=ValueError("corrupt template")):
            with self.assertRaisesRegex(ValueError, "corrupt template"):
                execute_prepare({"engineering_facts_path": "inputs/engineering_facts.json"}, content=self.content, host_client=self.host)

    def test_shared_resource_reference_validation_rejects_corruption(self) -> None:
        base_template = json.loads(TEMPLATE_PATH.read_text(encoding="utf-8"))
        base_rules = json.loads(CHAPTER_RULES_PATH.read_text(encoding="utf-8"))
        base_standards = json.loads(STANDARDS_LIBRARY_PATH.read_text(encoding="utf-8"))
        cases: list[tuple[str, dict[str, Any], dict[str, Any], dict[str, Any], str]] = []

        template = copy.deepcopy(base_template)
        next(block for section in template["sections"] if section["id"] == "1.1.1" for block in section["blocks"] if block["type"] == "deterministic_table")["builder"] = "unknown_builder"
        cases.append(("builder", template, copy.deepcopy(base_rules), copy.deepcopy(base_standards), "unknown deterministic builder"))
        template = copy.deepcopy(base_template)
        next(block for section in template["sections"] if section["id"] == "1.1.4" for block in section["blocks"] if block["type"] == "standards_reference")["group"] = "missing_group"
        cases.append(("standard", template, copy.deepcopy(base_rules), copy.deepcopy(base_standards), "unknown standards group"))
        template = copy.deepcopy(base_template)
        next(block for section in template["sections"] if section["id"] == "1.1.2" for block in section["blocks"] if block["type"] == "llm_section")["contract_ref"] = "missing_contract"
        cases.append(("contract", template, copy.deepcopy(base_rules), copy.deepcopy(base_standards), "must match template llm_section contract_ref"))
        rules = copy.deepcopy(base_rules)
        rules["writing"]["99.9"] = copy.deepcopy(rules["writing"]["1.1.2"])
        cases.append(("rule", copy.deepcopy(base_template), rules, copy.deepcopy(base_standards), "references missing template section"))

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            for name, template, rules, standards, expected in cases:
                with self.subTest(name=name):
                    paths = []
                    for suffix, payload in (("template", template), ("rules", rules), ("standards", standards)):
                        path = root / f"{name}_{suffix}.json"
                        path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
                        paths.append(path)
                    with self.assertRaisesRegex(ValueError, expected):
                        load_template(*paths)

    def test_exporter_failure_is_internal_exception(self) -> None:
        prepared, _ = self.prepare()
        with patch("src.report_finalize.core.export_report_files", side_effect=RuntimeError("export bug")):
            with self.assertRaisesRegex(RuntimeError, "export bug"):
                execute_finalize({"work_package_path": prepared["artifact"]["path"]}, content=self.content, host_client=self.host)

    def test_engineering_facts_core_does_not_import_report_shared(self) -> None:
        source = (TOOL_ROOT / "src" / "engineering_facts" / "core.py").read_text(encoding="utf-8")
        self.assertNotIn("src.report_shared", source)

    def test_generated_work_package_schema_failure_is_internal(self) -> None:
        with patch(
            "src.report_shared.workflow.validate_with_schema",
            side_effect=jsonschema.ValidationError("generated package invalid"),
        ):
            with self.assertRaisesRegex(jsonschema.ValidationError, "generated package invalid"):
                execute_prepare({"engineering_facts_path": "inputs/engineering_facts.json"}, content=self.content, host_client=self.host)

    def test_generated_manifest_schema_failure_is_internal_and_uncommitted(self) -> None:
        prepared, _ = self.prepare()
        with patch("src.report_finalize.core.build_manifest", side_effect=jsonschema.ValidationError("generated manifest invalid")):
            with self.assertRaisesRegex(jsonschema.ValidationError, "generated manifest invalid"):
                execute_finalize({"work_package_path": prepared["artifact"]["path"]}, content=self.content, host_client=self.host)
        self.assertFalse([path for path in self.host.files if path.endswith("/report_manifest.json")])


if __name__ == "__main__":
    unittest.main()
