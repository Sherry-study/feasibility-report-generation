from __future__ import annotations

import copy
import json
import os
import re
import tempfile
import unittest
from pathlib import Path
from typing import Any
from unittest.mock import patch
from urllib.parse import unquote, urlparse

import jsonschema
from docx import Document
from docx.oxml.ns import qn

from src.report_generation import execute
from src.report_generation.context_builder import FACT_ROOTS
from src.report_generation.template_loader import CHAPTER_RULES_PATH, STANDARDS_LIBRARY_PATH, load_template

TOOL_ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_PATH = TOOL_ROOT / "src" / "report_generation" / "templates" / "report_template.json"
PACKAGE_SCHEMA_PATH = TOOL_ROOT / "schemas" / "report_work_package.schema.json"
RESULTS_SCHEMA_PATH = TOOL_ROOT / "schemas" / "report_work_results.schema.json"
EXISTING_DRAFT_TITLES_PATH = (
    TOOL_ROOT / "tests" / "fixtures" / "existing_draft_section_titles.txt"
)
EXPECTED_WRITING_IDS = [
    "1.1.2",
    "1.1.3",
    "2.1.1",
    "2.1.2",
    "2.2",
    "2.4",
    "4.1.2",
    "4.1.3.1",
    "4.1.3.2",
    "4.1.3.3",
    "4.1.4",
    "4.2.1",
    "18.2",
]
EXPECTED_SYNTHESIS_IDS = [
    "1.2",
    "26.1.1",
    "26.1.2",
    "26.1.3",
    "26.2",
    "26.3",
    "26.4",
]
SUMMARY_FIELDS = [
    "conclusions",
    "key_facts",
    "conditions",
    "risks",
    "recommendations",
    "unresolved_items",
]


def minimal_facts() -> dict[str, Any]:
    return {
        "meta": {"schema_version": "2.0"},
        "sources": [],
        "basic_info": {"construction_unit": "测试建设单位"},
        "unit": {
            "name": "测试装置",
            "type": "联合生产装置",
            "annual_operating_hours": 8000,
            "design_case": {
                "feeds": [{"name": "原料A", "stream_id": "S1"}],
                "products": [
                    {
                        "name": "产品A",
                        "stream_id": "P1",
                        "annual_capacity": {"value": 10000, "unit": "吨/年"},
                        "hourly_rate": {"value": 1250, "unit": "kg/h"},
                        "is_main_product": True,
                    }
                ],
            },
            "retrofit_case": {
                "feeds": [{"name": "原料A", "stream_id": "S1"}],
                "products": [
                    {
                        "name": "产品A",
                        "stream_id": "P1",
                        "annual_capacity": {"value": 20000, "unit": "吨/年"},
                        "hourly_rate": {"value": 2500, "unit": "kg/h"},
                        "is_main_product": True,
                    }
                ],
            },
        },
        "component_catalog": [{"name": "产品A", "formula": "A"}],
        "process": {
            "design": {
                "streams": [
                    {
                        "stream_id": "S1",
                        "name": "原料A",
                        "source_equipment_id": "",
                        "target_equipment_id": "R1",
                        "flow": {"value": 1000, "unit": "kg/h"},
                    }
                ]
            },
            "retrofit": {
                "streams": [
                    {
                        "stream_id": "S1",
                        "name": "原料A",
                        "source_equipment_id": "",
                        "target_equipment_id": "R1",
                        "flow": {"value": 2000, "unit": "kg/h"},
                    }
                ]
            },
        },
        "diagnosis": {"bottlenecks": ["测试瓶颈"]},
        "scheme_analysis": {
            "candidates": [
                {
                    "name": "方案A",
                    "retrofit_type": "add_equipment",
                    "brief": "新增设备",
                    "evaluation": {
                        "technical_feasibility": {"level": "待核实"},
                        "implementation_complexity": {"level": "中"},
                        "operational_risk": {"level": "待核实"},
                        "energy_efficiency": {
                            "comment": "能耗影响待核实",
                            "display_name": "能效影响",
                        },
                    },
                }
            ],
            "user_candidates": [
                {
                    "name": "自定义方案",
                    "source_pool": "user_candidates",
                    "retrofit_type": "equipment_service_change",
                    "brief": "优化内件",
                    "evaluation": {
                        "maintainability": {"value": "中"},
                    },
                }
            ],
            "recommendation": {"name": "方案A"},
        },
        "adopted_scheme": {
            "selected_names": ["方案A"],
            "optimized": True,
            "schemes": [
                {
                    "name": "方案A",
                    "brief": "新增设备",
                    "description": "新增设备并调整局部流程连接。",
                }
            ],
        },
        "equipment": {
            "object_catalog": [
                {
                    "id": "1",
                    "name": "既有反应器",
                    "type": "reactor",
                    "tag": "R-101",
                    "specifications": {"反应类型": "测试反应"},
                },
                {
                    "id": "2",
                    "name": "既有精馏塔",
                    "type": "separator",
                    "specifications": {"类型": "一进两出型"},
                },
                {
                    "id": "2_1",
                    "name": "既有精馏塔_并联",
                    "type": "separator",
                    "is_new": True,
                    "specifications": {"类型": "新增并联塔"},
                },
                {
                    "id": "3",
                    "name": "新反应器",
                    "type": "reactor",
                    "is_new": True,
                    "specifications": {"反应类型": "利旧用途"},
                },
                {
                    "id": "4",
                    "name": "并联反应器",
                    "type": "reactor",
                    "specifications": {"反应类型": "并联用途"},
                },
            ],
            "reactor": {
                "selected_scheme": {
                    "scheme_id": "combined_1",
                    "scheme_index": 1,
                    "reuse": [{"id": "3", "name": "新反应器", "scheme_type": "reuse", "reused_from": "1"}],
                    "addition": [{"id": "1", "name": "既有反应器", "scheme_type": "increase_volume"}],
                    "parallel": [{"id": "4", "name": "并联反应器", "scheme_type": "add_parallel"}],
                },
                "combined_schemes": [
                    {
                        "scheme_id": "combined_1",
                        "scheme_index": 1,
                        "conclusion": "方案1：既有反应器扩容，并联反应器新增",
                    },
                    {
                        "scheme_id": "combined_2",
                        "scheme_index": 2,
                        "conclusion": "方案2：全部采用新增并联反应器",
                    },
                ],
                "evaluations": [
                    {
                        "id": "1",
                        "tag": "R-101",
                        "name": "既有反应器",
                        "form": "固定床",
                        "conclusion": "扩容后满足改造工况。",
                        "retrofit": {
                            "schemes": [
                                {
                                    "scheme_type": "increase_volume",
                                    "description": "扩容方案：更换更大规格反应器",
                                    "diameter": 1800,
                                    "cylinder_height": 12000,
                                    "total_volume": 40.0,
                                    "catalyst_volume": 24.0,
                                    "quantity": 1,
                                }
                            ]
                        },
                    },
                    {
                        "id": "3",
                        "tag": "",
                        "name": "新反应器",
                        "form": "固定床",
                        "retrofit": {
                            "schemes": [
                                {
                                    "scheme_type": "reuse",
                                    "description": "利旧方案：复用既有反应器",
                                    "diameter": 1600,
                                    "cylinder_height": 9000,
                                    "catalyst_volume": 18.0,
                                    "quantity": 1,
                                }
                            ]
                        },
                    },
                    {
                        "id": "4",
                        "tag": "R-104",
                        "name": "并联反应器",
                        "form": "列管式",
                        "retrofit": {
                            "schemes": [
                                {
                                    "scheme_type": "add_parallel",
                                    "description": "并联方案：增设并联反应器",
                                    "diameter": 1000,
                                    "cylinder_height": 6000,
                                    "catalyst_volume": 9.0,
                                    "heat_exchange_area": 80,
                                    "quantity": 2,
                                }
                            ]
                        },
                    },
                ],
            },
            "tower": {
                "combined_schemes": {
                    "optimal": [],
                    "parallel": [
                        [
                            {"id": "2", "name": "既有精馏塔"},
                            {"id": "2_1", "name": "既有精馏塔_并联", "is_new": True},
                        ]
                    ],
                    "series": [],
                    "other": [],
                },
                "evaluations": [
                    {
                        "id": "2",
                        "name": "既有精馏塔",
                        "retrofit": {
                            "route": "parallel",
                            "plan_detail": {
                                "device_paras": {
                                    "struct_info": [[800, 300], [1000, 450]],
                                    "tray_num": 24,
                                    "tray_type": "F1",
                                }
                            },
                        },
                    }
                ],
            },
        },
        "derived_facts": {
            "annualized_material_consumption": [
                {
                    "condition": "design",
                    "stream_id": "S1",
                    "stream_name": "原料A",
                    "annual_consumption": {"value": 8000, "unit": "t/a"},
                },
                {
                    "condition": "retrofit",
                    "stream_id": "S1",
                    "stream_name": "原料A",
                    "annual_consumption": {"value": 16000, "unit": "t/a"},
                },
            ],
            "utility_consumption_summary": {
                "status": "identified",
                "annual_operating_hours": 8000,
                "items": [
                    {
                        "status": "annualized",
                        "condition": "retrofit",
                        "equipment_id": "R-101",
                        "equipment_name": "新增反应器",
                        "medium": "electricity",
                        "medium_code": "electricity",
                        "medium_name": "电力",
                        "input": {"power_kw": 10, "annual_operating_hours": 8000},
                        "annual_quantity": {"value": 80000, "unit": "kWh/a"},
                        "formula": "power_kw * annual_operating_hours",
                    }
                ],
            },
            "energy_conversion": {
                "status": "calculated",
                "conversion_type": "standard_coal",
                "standard": "GB/T 2589-2020",
                "total_unit": "t标准煤（tce）",
                "items": [
                    {
                        "status": "calculated",
                        "medium": "electricity",
                        "medium_code": "electricity",
                        "medium_name": "电力",
                        "condition": "retrofit",
                        "equipment_id": "R-101",
                        "equipment_name": "新增反应器",
                        "quantity": 80000,
                        "unit": "kWh",
                        "rule_id": "SC-ELEC-001",
                        "coefficient": 0.1229,
                        "coefficient_unit": "kgce/kWh",
                        "standard_coal_tce": 9.832,
                        "source": {
                            "standard": "GB/T 2589-2020",
                            "rule_version": "0.3",
                        },
                    }
                ],
                "total_standard_coal_tce": 9.832,
            },
        },
    }


def uri_to_path(value: str) -> Path:
    parsed = urlparse(value)
    path = unquote(parsed.path)
    if len(path) >= 3 and path[0] == "/" and path[2] == ":":
        path = path[1:]
    return Path(path)


def summary() -> dict[str, list[str]]:
    return {field: [] for field in SUMMARY_FIELDS}


def implementation_schedule() -> dict[str, Any]:
    return {
        "rows": [
            {
                "stage": "项目前期（各报告编制及审批）",
                "main_content": "明确项目边界，完成可研及相关报告编制审批。",
                "duration_range": "2～3个月",
                "prerequisites": "采用方案和基础资料明确",
            },
            {
                "stage": "基础设计",
                "main_content": "完成工艺、设备、管道、自控等基础设计。",
                "duration_range": "2～3个月",
                "prerequisites": "可研阶段主要边界确认",
            },
            {
                "stage": "施工图设计",
                "main_content": "细化设备布置、管线改接、仪表联锁和施工图文件。",
                "duration_range": "2～3个月",
                "prerequisites": "基础设计审查完成",
            },
            {
                "stage": "设备采购",
                "main_content": "完成新增及改造设备采购、制造和到货验收。",
                "duration_range": "3～5个月",
                "prerequisites": "设备规格和请购条件明确",
            },
            {
                "stage": "土建施工",
                "main_content": "实施新增设备基础及相关土建配套。",
                "duration_range": "1～2个月",
                "prerequisites": "施工图和现场条件具备",
            },
            {
                "stage": "安装工程",
                "main_content": "完成设备安装、管道改接、电仪安装和系统联调。",
                "duration_range": "2～3个月",
                "prerequisites": "设备到货并具备停工施工窗口",
            },
            {
                "stage": "试生产",
                "main_content": "完成吹扫置换、单机试车、联动试车和投料试生产。",
                "duration_range": "1～2个月",
                "prerequisites": "安装质量验收和安全条件确认",
            },
        ]
    }


class ReportGenerationToolTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.root = Path(self.temp_dir.name)
        self.artifact_dir = self.root / "artifacts"
        self.facts_path = self.root / "engineering_facts.json"
        self.write_facts(minimal_facts())

    def write_facts(self, facts: dict[str, Any]) -> None:
        self.facts_path.write_text(
            json.dumps(facts, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

    def prepare(
        self,
        facts: dict[str, Any] | None = None,
        report_context: dict[str, Any] | None = None,
    ) -> tuple[dict[str, Any], dict[str, Any], Path]:
        if facts is not None:
            self.write_facts(facts)
        result = execute(
            {
                "operation": "prepare",
                "engineering_facts_uri": self.facts_path.as_uri(),
                "report_context": report_context or {},
            },
            self.artifact_dir,
        )
        package_path = self.artifact_dir / "work_package.json"
        package = (
            json.loads(package_path.read_text(encoding="utf-8"))
            if package_path.exists()
            else {}
        )
        return result, package, package_path

    def valid_results(self, package: dict[str, Any], package_path: Path) -> dict[str, Any]:
        return {
            "schema_version": "1.0",
            "writing_results": [
                {
                    "section_id": task["section_id"],
                    "title": "Agent 伪标题",
                    "table_id": "Agent 伪表号",
                    "blocks": [
                        {
                            "type": "paragraph",
                            "text": f"{task['section_id']} Agent 段落。",
                        }
                    ],
                    **(
                        {"implementation_schedule": implementation_schedule()}
                        if task["section_id"] == "18.2"
                        else {}
                    ),
                    "section_summary": summary(),
                }
                for task in package["writing_tasks"]
            ],
            "synthesis_results": [
                {
                    "section_id": task["section_id"],
                    "blocks": [
                        {
                            "type": "paragraph",
                            "text": f"{task['section_id']} 综合段落。",
                        }
                    ],
                    "section_summary": summary(),
                }
                for task in package["synthesis_tasks"]
            ],
        }

    def write_results(self, payload: dict[str, Any] | str) -> Path:
        results_path = self.root / "work_results.json"
        if isinstance(payload, str):
            results_path.write_text(payload, encoding="utf-8")
        else:
            results_path.write_text(
                json.dumps(payload, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
        return results_path

    def read_chapter_context(self, task: dict[str, Any]) -> dict[str, Any]:
        context_path = uri_to_path(task["context_uri"])
        self.assertTrue(context_path.is_file())
        return json.loads(context_path.read_text(encoding="utf-8"))

    def write_template_resources(
        self,
        template: dict[str, Any],
        chapter_rules: dict[str, Any],
        standards_library: dict[str, Any],
        name: str,
    ) -> tuple[Path, Path, Path]:
        template_path = self.root / f"{name}_template.json"
        chapter_rules_path = self.root / f"{name}_chapter_rules.json"
        standards_path = self.root / f"{name}_standards_library.json"
        template_path.write_text(json.dumps(template, ensure_ascii=False, indent=2), encoding="utf-8")
        chapter_rules_path.write_text(json.dumps(chapter_rules, ensure_ascii=False, indent=2), encoding="utf-8")
        standards_path.write_text(json.dumps(standards_library, ensure_ascii=False, indent=2), encoding="utf-8")
        return template_path, chapter_rules_path, standards_path

    def assert_prepare_fails_with_template_resources(
        self,
        template: dict[str, Any],
        chapter_rules: dict[str, Any],
        standards_library: dict[str, Any],
        name: str,
        expected_message: str,
    ) -> None:
        paths = self.write_template_resources(template, chapter_rules, standards_library, name)
        with patch(
            "src.report_generation.core.load_template",
            side_effect=lambda: load_template(*paths),
        ):
            result, package, _ = self.prepare()
        self.assertEqual(result["status"], "failed")
        self.assertEqual(package, {})
        self.assertIn(expected_message, result["diagnostics"][0]["message"])

    def finalize(
        self,
        package_path: Path,
        results_path: Path | None = None,
        artifact_dir: Path | None = None,
        work_results: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        request: dict[str, Any] = {
            "operation": "finalize",
            "work_package_uri": package_path.as_uri(),
        }
        if results_path is not None:
            request["work_results_uri"] = results_path.as_uri()
        if work_results is not None:
            request["work_results"] = work_results
        return execute(request, artifact_dir or self.artifact_dir)

    def test_prepare_generates_schema_valid_work_package(self) -> None:
        result, package, package_path = self.prepare()

        self.assertEqual(result["status"], "prepared")
        self.assertEqual(package["status"], "prepared")
        self.assertEqual(package["project_name"], "测试建设单位测试装置技术改造项目")
        self.assertEqual([s["order"] for s in package["sections"]], list(range(1, 67)))
        self.assertEqual(result["summary"]["writing_task_count"], 13)
        self.assertEqual(result["summary"]["synthesis_task_count"], 7)
        self.assertLessEqual(result["summary"]["research_task_count"], 4)
        self.assertLess(package_path.stat().st_size, 1_000_000)
        self.assertEqual([task["section_id"] for task in package["writing_tasks"]], EXPECTED_WRITING_IDS)
        self.assertEqual([task["section_id"] for task in package["synthesis_tasks"]], EXPECTED_SYNTHESIS_IDS)
        self.assertGreater(result["summary"]["deterministic_summary_count"], 0)
        self.assertEqual(result["summary"]["deterministic_summary_count"], len(package["deterministic_summaries"]))
        self.assertNotIn("fact_slice", json.dumps(package["writing_tasks"], ensure_ascii=False))
        self.assertNotIn("section_title", json.dumps(package["research_tasks"], ensure_ascii=False))
        section_ids = [section["id"] for section in package["sections"]]
        parent_ids = {
            section_id
            for section_id in section_ids
            for candidate in section_ids
            if candidate != section_id and candidate.startswith(f"{section_id}.")
        }
        blank_static_leaf_ids = [
            section["id"]
            for section in package["sections"]
            if section["id"] not in parent_ids
            and section["agent_slot"] is None
            and not section["deterministic_blocks"]
        ]
        self.assertEqual(blank_static_leaf_ids, [])

        schema = json.loads(PACKAGE_SCHEMA_PATH.read_text(encoding="utf-8"))
        jsonschema.validate(package, schema)

    def test_prepare_is_deterministic_for_identical_inputs(self) -> None:
        _, first_package, first_path = self.prepare()
        first_bytes = first_path.read_bytes()

        _, second_package, second_path = self.prepare()

        self.assertEqual(first_package, second_package)
        self.assertEqual(first_bytes, second_path.read_bytes())
        self.assertEqual(first_package["created_at"], "1970-01-01T00:00:00+00:00")

    def test_research_scope_is_narrow_and_optional_by_section(self) -> None:
        _, package, _ = self.prepare()

        research_by_group = {
            task["research_group"]: task["target_section_ids"]
            for task in package["research_tasks"]
        }
        self.assertLessEqual(len(research_by_group), 4)
        self.assertEqual(research_by_group["technology_research"], ["4.1.2"])

        writing_by_id = {
            task["section_id"]: task
            for task in package["writing_tasks"]
        }
        self.assertIsNone(writing_by_id["1.1.3"]["research_task_id"])
        self.assertIsNone(writing_by_id["4.1.3.2"]["research_task_id"])

    def test_synthesis_sources_are_single_round_and_closed(self) -> None:
        _, package, _ = self.prepare()

        writing_ids = {task["section_id"] for task in package["writing_tasks"]}
        synthesis_ids = {task["section_id"] for task in package["synthesis_tasks"]}
        deterministic_ids = set(package["deterministic_summaries"])
        for task in package["synthesis_tasks"]:
            for source_id in task["context"]["source_section_ids"]:
                self.assertNotIn(source_id, synthesis_ids)
                self.assertIn(source_id, writing_ids | deterministic_ids)

        referenced_deterministic_ids = {
            source_id
            for task in package["synthesis_tasks"]
            for source_id in task["context"]["source_section_ids"]
            if source_id not in writing_ids
        }
        self.assertEqual(referenced_deterministic_ids, deterministic_ids)
        for section_id, section_summary in package["deterministic_summaries"].items():
            self.assertEqual(set(section_summary), set(SUMMARY_FIELDS), section_id)
            self.assertTrue(any(section_summary.values()), section_id)

    def test_project_name_uses_four_level_priority(self) -> None:
        _, package, _ = self.prepare(report_context={"project_name": "显式项目名"})
        self.assertEqual(package["project_name"], "显式项目名")

        facts = minimal_facts()
        _, package, _ = self.prepare(facts)
        self.assertEqual(package["project_name"], "测试建设单位测试装置技术改造项目")

        facts = minimal_facts()
        facts["basic_info"]["construction_unit"] = ""
        _, package, _ = self.prepare(facts)
        self.assertEqual(package["project_name"], "测试装置技术改造项目")

        facts = minimal_facts()
        facts["basic_info"] = {}
        facts["unit"] = {}
        _, package, _ = self.prepare(facts)
        self.assertEqual(package["project_name"], "工业装置技术改造项目")

    def test_relative_artifact_dir_is_resolved_for_file_uris(self) -> None:
        relative_dir = Path("relative_report_output")
        old_cwd = Path.cwd()
        try:
            os.chdir(self.root)
            result = execute(
                {
                    "operation": "prepare",
                    "engineering_facts_uri": self.facts_path.as_uri(),
                    "report_context": {},
                },
                relative_dir,
            )
        finally:
            os.chdir(old_cwd)

        self.assertEqual(result["status"], "prepared")
        package_path = uri_to_path(result["artifact"]["uri"])
        self.assertTrue(package_path.is_absolute())
        self.assertTrue(package_path.is_file())

    def test_missing_top_level_fact_root_fails(self) -> None:
        facts = minimal_facts()
        facts.pop("equipment")

        result, package, _ = self.prepare(facts)

        self.assertEqual(result["status"], "failed")
        self.assertEqual(package, {})
        self.assertIn("missing required roots", result["diagnostics"][0]["message"])

    def test_root_internal_fields_missing_still_prepares(self) -> None:
        facts = {
            "meta": {"schema_version": "2.0"},
            "sources": [],
            "basic_info": {},
            "unit": {},
            "component_catalog": [],
            "process": {},
            "diagnosis": {},
            "scheme_analysis": {},
            "adopted_scheme": {},
            "equipment": {},
            "derived_facts": {},
        }

        result, package, _ = self.prepare(facts)

        self.assertEqual(result["status"], "prepared")
        self.assertEqual(package["engineering_facts"]["root_keys"], FACT_ROOTS)
        self.assertEqual(package["project_name"], "工业装置技术改造项目")

    def test_invalid_top_level_fact_contract_fails(self) -> None:
        cases = [
            ("extra_root", lambda facts: facts.update({"project": {}}), "unsupported roots"),
            ("bad_meta_version", lambda facts: facts["meta"].update({"schema_version": "1.0"}), "schema_version"),
            ("bad_root_type", lambda facts: facts.update({"sources": {}}), "invalid type"),
        ]
        for _, mutate, expected in cases:
            with self.subTest(expected=expected):
                facts = minimal_facts()
                mutate(facts)
                result, package, _ = self.prepare(facts)
                self.assertEqual(result["status"], "failed")
                self.assertEqual(package, {})
                self.assertIn(expected, result["diagnostics"][0]["message"])

    def test_fact_slices_are_isolated_by_section_contract(self) -> None:
        _, package, _ = self.prepare()

        task_112 = next(t for t in package["writing_tasks"] if t["section_id"] == "1.1.2")
        self.assertNotIn("fact_slice", task_112)
        context_112 = self.read_chapter_context(task_112)
        self.assertEqual(set(context_112), {"section_id", "fact_slice"})
        self.assertEqual(context_112["section_id"], "1.1.2")
        self.assertEqual(set(context_112["fact_slice"]), {"basic_info"})
        self.assertNotIn("equipment", context_112["fact_slice"])

        task_4132 = next(t for t in package["writing_tasks"] if t["section_id"] == "4.1.3.2")
        context_4132 = self.read_chapter_context(task_4132)
        self.assertEqual(
            set(context_4132["fact_slice"]),
            {"scheme_analysis"},
        )
        self.assertEqual(
            set(context_4132["fact_slice"]["scheme_analysis"]),
            {"candidates", "user_candidates"},
        )
        self.assertNotIn("process", context_4132["fact_slice"])
        self.assertNotIn("adopted_scheme", context_4132["fact_slice"])
        self.assertNotIn("recommendation", context_4132["fact_slice"]["scheme_analysis"])

        task_4131 = next(t for t in package["writing_tasks"] if t["section_id"] == "4.1.3.1")
        context_4131 = self.read_chapter_context(task_4131)
        self.assertEqual(set(context_4131["fact_slice"]["scheme_analysis"]), {"candidates", "user_candidates"})

        task_4133 = next(t for t in package["writing_tasks"] if t["section_id"] == "4.1.3.3")
        context_4133 = self.read_chapter_context(task_4133)
        self.assertEqual(set(context_4133["fact_slice"]), {"scheme_analysis", "adopted_scheme"})
        self.assertEqual(set(context_4133["fact_slice"]["scheme_analysis"]), {"recommendation"})
        self.assertTrue(context_4133["fact_slice"]["adopted_scheme"]["optimized"])
        self.assertIn("brief", context_4133["fact_slice"]["adopted_scheme"]["schemes"][0])
        self.assertIn("description", context_4133["fact_slice"]["adopted_scheme"]["schemes"][0])
        self.assertIn("不得在报告正文输出 optimized=true", "\n".join(task_4133["contract"]["prohibited_content"]))

        task_414 = next(t for t in package["writing_tasks"] if t["section_id"] == "4.1.4")
        context_414 = self.read_chapter_context(task_414)
        self.assertEqual(
            set(context_414["fact_slice"]),
            {"unit", "process", "diagnosis", "adopted_scheme", "equipment"},
        )
        self.assertEqual(set(context_414["fact_slice"]["equipment"]), {"reactor", "tower"})

        task_421 = next(t for t in package["writing_tasks"] if t["section_id"] == "4.2.1")
        context_421 = self.read_chapter_context(task_421)
        self.assertEqual(set(context_421["fact_slice"]), {"process", "adopted_scheme"})

        task_22 = next(t for t in package["writing_tasks"] if t["section_id"] == "2.2")
        context_22 = self.read_chapter_context(task_22)
        self.assertEqual(set(context_22["fact_slice"]), {"unit", "diagnosis", "adopted_scheme"})
        self.assertNotIn("process", context_22["fact_slice"])
        self.assertEqual(
            set(context_22["fact_slice"]["unit"]),
            {"design_case", "retrofit_case"},
        )

        task_412 = next(t for t in package["writing_tasks"] if t["section_id"] == "4.1.2")
        context_412 = self.read_chapter_context(task_412)
        self.assertEqual(set(context_412["fact_slice"]), {"unit"})
        self.assertEqual(
            set(context_412["fact_slice"]["unit"]),
            {"name", "type", "design_case", "retrofit_case"},
        )

        task_182 = next(t for t in package["writing_tasks"] if t["section_id"] == "18.2")
        context_182 = self.read_chapter_context(task_182)
        self.assertEqual(set(context_182["fact_slice"]), {"equipment", "adopted_scheme"})
        self.assertEqual(set(context_182["fact_slice"]["equipment"]), {"object_catalog"})

    def test_equipment_tables_use_structured_selected_actions(self) -> None:
        _, package, _ = self.prepare()
        section = next(section for section in package["sections"] if section["id"] == "4.2.4")
        tables = {
            block["caption"]: block
            for block in section["deterministic_blocks"]
            if block["type"] == "table"
        }

        self.assertEqual(
            set(tables),
            {"新增工艺设备汇总表", "利旧工艺设备汇总表", "改造工艺设备汇总表"},
        )
        self.assertEqual(len(tables["新增工艺设备汇总表"]["rows"]), 2)
        self.assertEqual(len(tables["利旧工艺设备汇总表"]["rows"]), 1)
        self.assertEqual(len(tables["改造工艺设备汇总表"]["rows"]), 1)
        self.assertEqual(
            tables["改造工艺设备汇总表"]["headers"],
            ["序号", "设备类别", "设备位号", "设备名称", "主要改造内容", "改造后主要规格/设计参数", "数量", "备注"],
        )

        serialized = json.dumps(section, ensure_ascii=False)
        self.assertIn("既有精馏塔（并联新增）", serialized)
        self.assertIn("并联反应器（并联新增）", serialized)
        self.assertIn("反应器（利旧设备）", serialized)
        self.assertIn("扩容方案：更换更大规格反应器", serialized)
        self.assertIn("直径800/1000 mm", serialized)
        self.assertNotIn("新增设备\", \"待补充", serialized)
        self.assertNotIn("利旧设备\", \"待补充", serialized)
        self.assertNotIn("改造设备\", \"待补充", serialized)

    def test_scheme_comparison_table_uses_dynamic_evaluation_union(self) -> None:
        _, package, _ = self.prepare()
        section = next(section for section in package["sections"] if section["id"] == "4.1.3.2")
        table = next(
            block
            for block in section["deterministic_blocks"]
            if block["type"] == "table" and block["caption"] == "候选工艺技术方案比选表"
        )

        self.assertEqual(
            table["headers"],
            [
                "方案",
                "来源",
                "改造类型",
                "方案概要",
                "技术可行性",
                "实施复杂度",
                "运行风险",
                "能效影响",
                "Maintainability",
            ],
        )
        self.assertEqual(table["rows"][0][7], "能耗影响待核实")
        self.assertEqual(table["rows"][0][8], "待补充")
        self.assertEqual(table["rows"][1][4:8], ["待补充", "待补充", "待补充", "待补充"])
        self.assertEqual(table["rows"][1][8], "中")

    def test_key_equipment_comparison_table_marks_selected_and_alternative(self) -> None:
        _, package, _ = self.prepare()
        section = next(section for section in package["sections"] if section["id"] == "4.1.3.2")
        table = next(
            block
            for block in section["deterministic_blocks"]
            if block["type"] == "table" and block["caption"] == "关键设备方案比较表"
        )

        self.assertEqual(table["table_id"], "表4.2")
        self.assertEqual(table["headers"], ["设备名称", "候选路线", "推荐结论"])
        self.assertEqual(
            table["rows"],
            [
                ["反应器组", "方案1：既有反应器扩容，并联反应器新增", "选定"],
                ["反应器组", "方案2：全部采用新增并联反应器", "备选"],
            ],
        )

    def test_key_equipment_comparison_skips_with_warning_on_bad_selection(self) -> None:
        facts = minimal_facts()
        facts["equipment"]["reactor"]["selected_scheme"]["scheme_id"] = "missing"
        facts["equipment"]["reactor"]["selected_scheme"]["scheme_index"] = 99

        result, package, _ = self.prepare(facts)

        self.assertEqual(result["status"], "prepared")
        self.assertIn(
            "KEY_EQUIPMENT_COMPARISON_SKIPPED",
            {item["code"] for item in result["diagnostics"]},
        )
        section = next(section for section in package["sections"] if section["id"] == "4.1.3.2")
        captions = [
            block["caption"]
            for block in section["deterministic_blocks"]
            if block["type"] == "table"
        ]
        self.assertNotIn("关键设备方案比较表", captions)

    def test_synthesis_tasks_do_not_include_full_section_text(self) -> None:
        _, package, _ = self.prepare()

        task = next(t for t in package["synthesis_tasks"] if t["section_id"] == "26.2")
        self.assertEqual(
            set(task["context"]),
            {"source_section_ids", "summary_fields"},
        )
        serialized = json.dumps(task, ensure_ascii=False)
        self.assertNotIn("fact_slice", serialized)
        self.assertNotIn("deterministic_blocks", serialized)
        self.assertNotIn("full_text", serialized)

    def test_resource_reference_errors_fail_prepare(self) -> None:
        base_template = json.loads(TEMPLATE_PATH.read_text(encoding="utf-8"))
        base_rules = json.loads(CHAPTER_RULES_PATH.read_text(encoding="utf-8"))
        base_standards = json.loads(STANDARDS_LIBRARY_PATH.read_text(encoding="utf-8"))

        template = copy.deepcopy(base_template)
        block = next(
            block
            for section in template["sections"]
            if section["id"] == "1.1.1"
            for block in section["blocks"]
            if block["type"] == "deterministic_table"
        )
        block["builder"] = "unknown_builder"
        self.assert_prepare_fails_with_template_resources(
            template,
            copy.deepcopy(base_rules),
            copy.deepcopy(base_standards),
            "unknown_builder",
            "unknown deterministic builder",
        )

        template = copy.deepcopy(base_template)
        block = next(
            block
            for section in template["sections"]
            if section["id"] == "1.1.4"
            for block in section["blocks"]
            if block["type"] == "standards_reference"
        )
        block["group"] = "missing_group"
        self.assert_prepare_fails_with_template_resources(
            template,
            copy.deepcopy(base_rules),
            copy.deepcopy(base_standards),
            "unknown_standards_group",
            "unknown standards group",
        )

        template = copy.deepcopy(base_template)
        block = next(
            block
            for section in template["sections"]
            if section["id"] == "1.1.2"
            for block in section["blocks"]
            if block["type"] == "llm_section"
        )
        block["contract_ref"] = "missing_contract"
        self.assert_prepare_fails_with_template_resources(
            template,
            copy.deepcopy(base_rules),
            copy.deepcopy(base_standards),
            "bad_contract_ref",
            "must match template llm_section contract_ref",
        )

        rules = copy.deepcopy(base_rules)
        rules["writing"]["99.9"] = copy.deepcopy(rules["writing"]["1.1.2"])
        self.assert_prepare_fails_with_template_resources(
            copy.deepcopy(base_template),
            rules,
            copy.deepcopy(base_standards),
            "missing_rule_section",
            "references missing template section",
        )

    def test_request_rejects_undeclared_fields(self) -> None:
        prepare_result = execute(
            {
                "operation": "prepare",
                "engineering_facts_uri": self.facts_path.as_uri(),
                "report_context": {},
                "project_type": "不允许",
            },
            self.artifact_dir,
        )
        self.assertEqual(prepare_result["status"], "failed")
        self.assertIn("unsupported fields", prepare_result["diagnostics"][0]["message"])

        report_context_result = execute(
            {
                "operation": "prepare",
                "engineering_facts_uri": self.facts_path.as_uri(),
                "report_context": {"project_name": "显式项目名", "project_type": "不允许"},
            },
            self.artifact_dir,
        )
        self.assertEqual(report_context_result["status"], "failed")
        self.assertIn("report_context has unsupported fields", report_context_result["diagnostics"][0]["message"])

        finalize_result = execute(
            {
                "operation": "finalize",
                "work_package_uri": "file:///tmp/work_package.json",
                "work_results_uri": "file:///tmp/work_results.json",
                "repair": True,
            },
            self.artifact_dir,
        )
        self.assertEqual(finalize_result["status"], "failed")
        self.assertIn("unsupported fields", finalize_result["diagnostics"][0]["message"])

    def test_finalize_exports_markdown_and_docx(self) -> None:
        _, package, package_path = self.prepare()
        results = self.valid_results(package, package_path)
        results_path = self.write_results(results)

        result = self.finalize(package_path, results_path)

        self.assertEqual(result["status"], "completed")
        self.assertEqual(result["summary"]["fallback_section_count"], 0)
        markdown_path = self.artifact_dir / "可行性研究报告_初稿.md"
        docx_path = self.artifact_dir / "可行性研究报告_初稿.docx"
        self.assertTrue(markdown_path.is_file())
        self.assertTrue(docx_path.is_file())
        markdown = markdown_path.read_text(encoding="utf-8")
        self.assertIn("测试建设单位测试装置技术改造项目", markdown)
        self.assertIn("**表18.1 项目实施进度计划表**", markdown)
        self.assertIn("| 阶段 | 主要内容 | 预计时长 | 前置条件 |", markdown)
        self.assertIn("项目前期（各报告编制及审批）", markdown)

        doc = Document(docx_path)
        schedule_tables = [
            table
            for table in doc.tables
            if [cell.text for cell in table.rows[0].cells]
            == ["阶段", "主要内容", "预计时长", "前置条件"]
        ]
        self.assertEqual(len(schedule_tables), 1)
        self.assertEqual(
            [cell.text for cell in schedule_tables[0].rows[1].cells],
            [
                "项目前期（各报告编制及审批）",
                "明确项目边界，完成可研及相关报告编制审批。",
                "2～3个月",
                "采用方案和基础资料明确",
            ],
        )
        self.assertTrue(
            all(
                row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None
                for row in schedule_tables[0].rows
            )
        )

    def test_finalize_without_agent_results_generates_incomplete_report(self) -> None:
        _, package, package_path = self.prepare()

        result = self.finalize(package_path)

        self.assertEqual(result["status"], "completed")
        self.assertEqual(
            result["summary"]["fallback_section_count"],
            len(package["writing_tasks"]) + len(package["synthesis_tasks"]),
        )
        codes = [item["code"] for item in result["diagnostics"]]
        self.assertEqual(codes[0], "WORK_RESULTS_NOT_PROVIDED")
        self.assertEqual(codes.count("SECTION_FALLBACK"), result["summary"]["fallback_section_count"])
        warning = result["diagnostics"][0]
        self.assertEqual(warning["level"], "warning")
        self.assertIn("不完整初稿", warning["message"])
        markdown = (self.artifact_dir / "可行性研究报告_初稿.md").read_text(encoding="utf-8")
        agent_sections = [
            section for section in package["sections"] if section["agent_slot"] is not None
        ]
        for section in agent_sections:
            self.assertIn(section["fallback"], markdown)

    def test_finalize_with_inline_work_results(self) -> None:
        _, package, package_path = self.prepare()
        results = self.valid_results(package, package_path)

        result = self.finalize(package_path, work_results=results)

        self.assertEqual(result["status"], "completed")
        self.assertEqual(result["summary"]["fallback_section_count"], 0)
        self.assertNotIn("WORK_RESULTS_NOT_PROVIDED", {item["code"] for item in result["diagnostics"]})
        markdown = (self.artifact_dir / "可行性研究报告_初稿.md").read_text(encoding="utf-8")
        self.assertIn("1.1.2 Agent 段落。", markdown)
        docx_path = self.artifact_dir / "可行性研究报告_初稿.docx"
        self.assertTrue(docx_path.is_file())

    def test_finalize_rejects_work_results_uri_and_inline_conflict(self) -> None:
        _, package, package_path = self.prepare()
        results = self.valid_results(package, package_path)
        results_path = self.write_results(results)

        result = self.finalize(package_path, results_path, work_results=results)

        self.assertEqual(result["status"], "failed")
        self.assertIn("mutually exclusive", result["diagnostics"][0]["message"])

    def test_finalize_with_invalid_inline_work_results_fails(self) -> None:
        _, _, package_path = self.prepare()

        incomplete = self.finalize(
            package_path,
            work_results={"schema_version": "1.0", "writing_results": []},
        )
        self.assertEqual(incomplete["status"], "failed")
        self.assertIn("missing required fields", incomplete["diagnostics"][0]["message"])

        not_object = execute(
            {
                "operation": "finalize",
                "work_package_uri": package_path.as_uri(),
                "work_results": ["not", "an", "object"],
            },
            self.artifact_dir,
        )
        self.assertEqual(not_object["status"], "failed")
        self.assertIn("work_results must be an object", not_object["diagnostics"][0]["message"])

    def test_finalize_with_blank_work_results_uri_fails(self) -> None:
        _, _, package_path = self.prepare()

        result = execute(
            {
                "operation": "finalize",
                "work_package_uri": package_path.as_uri(),
                "work_results_uri": "   ",
            },
            self.artifact_dir,
        )

        self.assertEqual(result["status"], "failed")
        self.assertIn("work_results_uri", result["diagnostics"][0]["message"])

    def test_missing_or_invalid_dynamic_sections_use_fallback(self) -> None:
        _, package, package_path = self.prepare()
        results = self.valid_results(package, package_path)
        results["writing_results"] = [
            item for item in results["writing_results"] if item["section_id"] != "1.1.2"
        ]
        invalid = next(item for item in results["writing_results"] if item["section_id"] == "1.1.3")
        invalid["blocks"] = [{"type": "table", "text": "Agent 非法块"}]
        results_path = self.write_results(results)

        result = self.finalize(package_path, results_path)

        self.assertEqual(result["status"], "completed")
        self.assertEqual(result["summary"]["fallback_section_count"], 2)
        self.assertEqual(
            [item["code"] for item in result["diagnostics"]],
            ["SECTION_FALLBACK", "SECTION_FALLBACK"],
        )
        markdown = (self.artifact_dir / "可行性研究报告_初稿.md").read_text(encoding="utf-8")
        section_112 = next(s for s in package["sections"] if s["id"] == "1.1.2")
        section_113 = next(s for s in package["sections"] if s["id"] == "1.1.3")
        self.assertIn(section_112["fallback"], markdown)
        self.assertIn(section_113["fallback"], markdown)

    def test_empty_dynamic_blocks_use_fallback(self) -> None:
        _, package, package_path = self.prepare()
        results = self.valid_results(package, package_path)
        target = next(
            item for item in results["writing_results"]
            if item["section_id"] == "1.1.2"
        )
        target["blocks"] = []
        results_path = self.write_results(results)

        result = self.finalize(package_path, results_path)

        self.assertEqual(result["status"], "completed")
        self.assertEqual(result["summary"]["fallback_section_count"], 1)

    def test_implementation_schedule_invalid_variants_use_18_2_fallback(self) -> None:
        cases = {
            "missing_schedule": lambda schedule: None,
            "missing_row": lambda schedule: {**schedule, "rows": schedule["rows"][:-1]},
            "duplicate_stage": lambda schedule: {
                **schedule,
                "rows": [
                    schedule["rows"][0],
                    {**schedule["rows"][0]},
                    *schedule["rows"][2:],
                ],
            },
            "wrong_order": lambda schedule: {
                **schedule,
                "rows": [schedule["rows"][1], schedule["rows"][0], *schedule["rows"][2:]],
            },
            "extra_row": lambda schedule: {
                **schedule,
                "rows": [*schedule["rows"], {**schedule["rows"][-1], "stage": "额外阶段"}],
            },
            "bad_duration": lambda schedule: {
                **schedule,
                "rows": [
                    {**schedule["rows"][0], "duration_range": "2个月"},
                    *schedule["rows"][1:],
                ],
            },
        }
        for name, mutate in cases.items():
            with self.subTest(name=name):
                artifact_dir = self.root / f"artifacts_{name}"
                _, package, package_path = self.prepare()
                results = self.valid_results(package, package_path)
                target = next(item for item in results["writing_results"] if item["section_id"] == "18.2")
                mutated = mutate(copy.deepcopy(target["implementation_schedule"]))
                if mutated is None:
                    target.pop("implementation_schedule")
                else:
                    target["implementation_schedule"] = mutated
                results_path = self.write_results(results)

                result = self.finalize(package_path, results_path, artifact_dir)

                self.assertEqual(result["status"], "completed")
                self.assertEqual(result["summary"]["fallback_section_count"], 1)
                self.assertIn("SECTION_FALLBACK", {item["code"] for item in result["diagnostics"]})
                markdown = (artifact_dir / "可行性研究报告_初稿.md").read_text(encoding="utf-8")
                section_182 = next(s for s in package["sections"] if s["id"] == "18.2")
                self.assertIn(section_182["fallback"], markdown)
                self.assertNotIn("**表18.1 项目实施进度计划表**", markdown)

    def test_implementation_schedule_is_rejected_outside_18_2(self) -> None:
        _, package, package_path = self.prepare()
        results = self.valid_results(package, package_path)
        target = next(item for item in results["writing_results"] if item["section_id"] == "1.1.2")
        target["implementation_schedule"] = implementation_schedule()
        results_path = self.write_results(results)

        result = self.finalize(package_path, results_path)

        self.assertEqual(result["status"], "failed")
        self.assertIn("implementation_schedule", result["diagnostics"][0]["message"])

    def test_agent_output_must_not_expose_raw_optimized_flags(self) -> None:
        _, package, package_path = self.prepare()
        results = self.valid_results(package, package_path)
        target = next(
            item for item in results["writing_results"]
            if item["section_id"] == "4.1.3.3"
        )
        target["blocks"] = [{"type": "paragraph", "text": "最终采用方案B，optimized=true。"}]
        results_path = self.write_results(results)

        result = self.finalize(package_path, results_path)

        self.assertEqual(result["status"], "completed")
        self.assertEqual(result["summary"]["fallback_section_count"], 1)
        markdown = (self.artifact_dir / "可行性研究报告_初稿.md").read_text(encoding="utf-8")
        self.assertNotIn("optimized=true", markdown)

    def test_missing_result_collection_fails(self) -> None:
        _, package, package_path = self.prepare()
        results = self.valid_results(package, package_path)
        del results["synthesis_results"]
        results_path = self.write_results(results)

        result = self.finalize(package_path, results_path)

        self.assertEqual(result["status"], "failed")
        self.assertIn("missing required fields", result["diagnostics"][0]["message"])

    def test_duplicate_or_unknown_result_section_id_fails(self) -> None:
        _, package, package_path = self.prepare()
        results = self.valid_results(package, package_path)
        results["writing_results"].append(copy.deepcopy(results["writing_results"][0]))
        duplicate_path = self.write_results(results)

        duplicate_result = self.finalize(package_path, duplicate_path)

        self.assertEqual(duplicate_result["status"], "failed")
        self.assertIn("duplicate section_id", duplicate_result["diagnostics"][0]["message"])

        results = self.valid_results(package, package_path)
        results["writing_results"][0]["section_id"] = "unknown"
        unknown_path = self.write_results(results)

        unknown_result = self.finalize(package_path, unknown_path)

        self.assertEqual(unknown_result["status"], "failed")
        self.assertIn("unknown section_id", unknown_result["diagnostics"][0]["message"])

    def test_agent_cannot_control_titles_or_table_numbers(self) -> None:
        _, package, package_path = self.prepare()
        results = self.valid_results(package, package_path)
        results_path = self.write_results(results)

        result = self.finalize(package_path, results_path)

        self.assertEqual(result["status"], "completed")
        markdown = (self.artifact_dir / "可行性研究报告_初稿.md").read_text(encoding="utf-8")
        self.assertIn("### 1.1.2 主办单位基本情况", markdown)
        self.assertNotIn("Agent 伪标题", markdown)
        self.assertNotIn("Agent 伪表号", markdown)

    def test_agent_paragraph_bullet_and_numbered_blocks_export_to_both_formats(self) -> None:
        _, package, package_path = self.prepare()
        results = self.valid_results(package, package_path)
        target = next(item for item in results["writing_results"] if item["section_id"] == "1.1.2")
        target["blocks"] = [
            {"type": "paragraph", "text": "段落块导出检查。"},
            {"type": "bullet_list", "items": ["无序项A", "无序项B"]},
            {"type": "numbered_list", "items": ["有序项A", "有序项B"]},
        ]
        results_path = self.write_results(results)

        result = self.finalize(package_path, results_path)

        self.assertEqual(result["status"], "completed")
        markdown = (self.artifact_dir / "可行性研究报告_初稿.md").read_text(encoding="utf-8")
        self.assertIn("段落块导出检查。", markdown)
        self.assertIn("- 无序项A", markdown)
        self.assertIn("1. 有序项A", markdown)

        doc = Document(self.artifact_dir / "可行性研究报告_初稿.docx")
        paragraphs = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("段落块导出检查。", paragraphs)
        self.assertIn("无序项A", paragraphs)
        self.assertIn("有序项A", paragraphs)

    def test_bad_json_fails(self) -> None:
        _, package, package_path = self.prepare()
        results_path = self.write_results("{bad json")

        result = self.finalize(package_path, results_path)

        self.assertEqual(result["status"], "failed")
        self.assertIn("REPORT_TOOL_FAILED", {item["code"] for item in result["diagnostics"]})

    def test_work_results_schema_accepts_agent_metadata_but_block_shape_is_fixed(self) -> None:
        _, package, package_path = self.prepare()
        results = self.valid_results(package, package_path)
        schema = json.loads(RESULTS_SCHEMA_PATH.read_text(encoding="utf-8"))

        jsonschema.validate(results, schema)

    def test_template_covers_existing_draft_sections_and_fallbacks_are_bounded(self) -> None:
        if not EXISTING_DRAFT_TITLES_PATH.is_file():
            self.skipTest("existing draft section-title fixture is not available")
        template = json.loads(TEMPLATE_PATH.read_text(encoding="utf-8"))
        template_titles = {section["title"] for section in template["sections"]}
        draft_titles = [
            line.strip()
            for line in EXISTING_DRAFT_TITLES_PATH.read_text(encoding="utf-8").splitlines()
            if line.strip()
        ]

        self.assertFalse([title for title in draft_titles if title not in template_titles])
        self.assertFalse([section["id"] for section in template["sections"] if "order" in section])
        self.assertFalse([section["id"] for section in template["sections"] if "contract" in section])
        embedded_standards = [
            section["id"]
            for section in template["sections"]
            for block in section.get("blocks", [])
            if block.get("type") == "standards_reference" and "items" in block
        ]
        self.assertEqual(embedded_standards, [])
        forbidden = re.compile(r"满足规范|依托现有系统|依托现有罐区|依托既有|依托现有管网")
        bad_fallbacks = [
            section["id"]
            for section in template["sections"]
            if forbidden.search(section.get("fallback", ""))
        ]
        self.assertEqual(bad_fallbacks, [])

    def test_utilities_and_energy_tables_are_dynamic(self) -> None:
        _, package, _ = self.prepare()
        utilities_section = next(section for section in package["sections"] if section["id"] == "5.4")
        energy_section = next(section for section in package["sections"] if section["id"] == "10.5")
        utilities_text = json.dumps(utilities_section, ensure_ascii=False)
        energy_text = json.dumps(energy_section, ensure_ascii=False)
        self.assertIn("新增反应器", utilities_text)
        self.assertIn("9.83", energy_text)
        self.assertIn("SC-ELEC-001", energy_text)
        self.assertNotIn("待核实", energy_text)

    def test_report_tables_keep_missing_conversion_reasons(self) -> None:
        facts = minimal_facts()
        facts["derived_facts"]["utility_consumption_summary"]["items"] = [
            {
                "status": "annualized",
                "condition": "retrofit",
                "equipment_name": "冷却设备",
                "medium": "CW",
                "medium_code": "cooling_water",
                "medium_name": "循环冷却水",
                "annual_quantity": {"value": 1200, "unit": "t/a"},
                "formula": "mass_flow_kg_h * annual_operating_hours / 1000",
            }
        ]
        facts["derived_facts"]["energy_conversion"] = {
            "status": "blocked",
            "conversion_type": "standard_coal",
            "items": [
                {
                    "status": "blocked",
                    "reason": "conversion_factor_missing",
                    "medium": "CW",
                    "medium_code": "cooling_water",
                    "medium_name": "循环冷却水",
                    "quantity": 1200,
                    "unit": "t",
                }
            ],
        }

        _, package, _ = self.prepare(facts)
        serialized = json.dumps(package["sections"], ensure_ascii=False)

        self.assertIn("循环冷却水", serialized)
        self.assertIn("缺少折标系数", serialized)
        self.assertNotIn("conversion_factor_missing", serialized)
        self.assertNotIn("导热油", serialized)

    def test_blocked_energy_table_does_not_show_zero_total(self) -> None:
        facts = minimal_facts()
        facts["derived_facts"]["energy_conversion"] = {
            "status": "blocked",
            "conversion_type": "standard_coal",
            "items": [
                {
                    "status": "blocked",
                    "reason": "conversion_factor_missing",
                    "medium": "CW",
                    "medium_code": "cooling_water",
                    "medium_name": "循环冷却水",
                    "quantity": 1200,
                    "unit": "t",
                }
            ],
        }

        _, package, _ = self.prepare(facts)
        energy_section = next(section for section in package["sections"] if section["id"] == "10.5")
        table = energy_section["deterministic_blocks"][0]

        self.assertEqual(len(table["rows"]), 1)
        self.assertNotIn("合计", json.dumps(table, ensure_ascii=False))
        self.assertNotIn("0.00", json.dumps(table, ensure_ascii=False))

    def test_energy_table_separates_condition_subtotals(self) -> None:
        facts = minimal_facts()
        facts["derived_facts"]["energy_conversion"] = {
            "status": "calculated",
            "conversion_type": "standard_coal",
            "items": [
                {
                    "status": "calculated",
                    "medium": "electricity",
                    "medium_code": "electricity",
                    "medium_name": "电力",
                    "condition": "design",
                    "quantity": 80000,
                    "unit": "kWh",
                    "rule_id": "SC-ELEC-001",
                    "coefficient": 0.1229,
                    "coefficient_unit": "kgce/kWh",
                    "standard_coal_tce": 9.832,
                },
                {
                    "status": "calculated",
                    "medium": "electricity",
                    "medium_code": "electricity",
                    "medium_name": "电力",
                    "condition": "retrofit",
                    "quantity": 160000,
                    "unit": "kWh",
                    "rule_id": "SC-ELEC-001",
                    "coefficient": 0.1229,
                    "coefficient_unit": "kgce/kWh",
                    "standard_coal_tce": 19.664,
                },
            ],
            "totals_by_condition": [
                {"condition": "design", "item_count": 1, "total_standard_coal_tce": 9.832},
                {"condition": "retrofit", "item_count": 1, "total_standard_coal_tce": 19.664},
            ],
        }

        _, package, _ = self.prepare(facts)
        energy_section = next(section for section in package["sections"] if section["id"] == "10.5")
        table = energy_section["deterministic_blocks"][0]
        serialized = json.dumps(table, ensure_ascii=False)

        self.assertIn("工况", table["headers"])
        self.assertIn("改造前", serialized)
        self.assertIn("改造后", serialized)
        self.assertEqual(
            [row for row in table["rows"] if row[0] == "小计"],
            [
                ["小计", "改造前", "", "", "", "9.83", "已折算"],
                ["小计", "改造后", "", "", "", "19.66", "已折算"],
            ],
        )
        self.assertNotIn("29.50", serialized)

    def test_no_utility_data_uses_single_gap_row(self) -> None:
        facts = minimal_facts()
        facts["derived_facts"].pop("utility_consumption_summary")
        facts["derived_facts"].pop("energy_conversion")

        _, package, _ = self.prepare(facts)
        utilities_section = next(section for section in package["sections"] if section["id"] == "5.4")
        energy_section = next(section for section in package["sections"] if section["id"] == "10.5")
        serialized = json.dumps([utilities_section, energy_section], ensure_ascii=False)

        self.assertIn("未识别到可汇总公用工程/待补充", serialized)
        self.assertNotIn("蒸汽", serialized)
        self.assertNotIn("循环冷却水", serialized)


if __name__ == "__main__":
    unittest.main()
