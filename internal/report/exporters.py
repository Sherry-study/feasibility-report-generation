from __future__ import annotations
import re
from pathlib import Path
from docx import Document
from docx.shared import Mm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# ===== 排版规范（report_formatter 契约；未定义项按既有级别规律推测） =====
# 页面: A4 纵向, 边距 上25/下25/左31.7/右31.7 mm -> 可用宽度约 8300 twips
# 正文: 中文宋体, 西文 Times New Roman, 10.5pt, 1.5倍行距, 两端对齐
# 标题: 一级18pt(编号"1.", 段前分页) / 二级16pt / 三级16pt / 四级14pt(推测: 按递减规律)
CN_FONT='宋体'
EN_FONT='Times New Roman'
BODY_SIZE=10.5
HEADING_SIZES={1:18,2:16,3:16,4:14}
USABLE_TWIPS=8300

def _set_east_asia(rPr, cn=CN_FONT, en=EN_FONT):
    rFonts=rPr.first_child_found_in('w:rFonts')
    if rFonts is None:
        rFonts=OxmlElement('w:rFonts'); rPr.append(rFonts)
    # 模板样式的 rFonts 带主题字体属性(asciiTheme/eastAsiaTheme 等), Word 解析时优先于
    # 显式 w:ascii/w:eastAsia, 不移除会导致标题渲染为等线/Calibri Light 而非宋体/Times New Roman
    for attr in ('asciiTheme','eastAsiaTheme','hAnsiTheme','cstheme'):
        rFonts.attrib.pop(qn('w:'+attr),None)
    rFonts.set(qn('w:ascii'),en); rFonts.set(qn('w:hAnsi'),en); rFonts.set(qn('w:eastAsia'),cn)

def set_font(run,size=None,bold=None):
    run.font.name=EN_FONT
    _set_east_asia(run._element.get_or_add_rPr())
    if size is not None: run.font.size=Pt(size)
    if bold is not None: run.bold=bold

def repeat_header(row):
    trPr=row._tr.get_or_add_trPr(); el=OxmlElement('w:tblHeader'); el.set(qn('w:val'),'true'); trPr.append(el)

def cell_width(cell,twips):
    tcPr=cell._tc.get_or_add_tcPr(); tcW=tcPr.first_child_found_in('w:tcW')
    if tcW is None: tcW=OxmlElement('w:tcW'); tcPr.append(tcW)
    tcW.set(qn('w:w'),str(twips)); tcW.set(qn('w:type'),'dxa')

def _table_font_size(ncols):
    """配置基准 10.5pt；列数多时按阶梯缩小，避免超出可用页宽（推测项）。"""
    if ncols<=5: return 10.5
    if ncols==6: return 9.5
    return 9.0

def set_cell_text(cell,text,bold=False,size=BODY_SIZE):
    cell.text=''; p=cell.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.5
    r=p.add_run('' if text is None else str(text)); set_font(r,size,bold); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def widths_for(n,table_id=None):
    specs={
      '3.1-1':[2100,1100,1500,1500,3000], '4.2-1':[1200,1700,900,900,1300,1300,2000],
      '4.2-3':[650,1100,1250,1550,2850,650,1950], '4.2-4':[650,1100,1250,1550,2850,650,1950],
      '4.2-5':[550,950,1050,1350,2300,2350,550,1200], '5.1-1':[1500,1700,850,1100,1100,1300,1800,1300],
      '5.4-1':[1400,1500,850,1000,1300,1700,2200], '8.1-1':[1400,1500,850,1000,1300,1700,2200],
      '10.5-1':[1900,1100,1700,1900,1500], '4.1-2':[1000,2300,850,850,900,900,900,850], '18.2-1':[1800,2200,1700,2200,1700], '19-1':[1900,1400,3300,2400]
    }
    if table_id in specs:
        w=specs[table_id]; scale=USABLE_TWIPS/sum(w)
        return [int(x*scale) for x in w]
    return [USABLE_TWIPS//max(n,1)]*n

def configure(doc):
    sec=doc.sections[0]
    sec.page_width=Mm(210); sec.page_height=Mm(297)
    sec.top_margin=Mm(25); sec.bottom_margin=Mm(25); sec.left_margin=Mm(31.7); sec.right_margin=Mm(31.7)
    styles=doc.styles
    normal=styles['Normal']; normal.font.name=EN_FONT
    _set_east_asia(normal.element.get_or_add_rPr())
    normal.font.size=Pt(BODY_SIZE)
    for name,size in [('Title',22),('Heading 1',18),('Heading 2',16),('Heading 3',16),('Heading 4',14)]:
        s=styles[name]; s.font.name=EN_FONT
        _set_east_asia(s.element.get_or_add_rPr())
        s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor(0,0,0)
    return doc

def normalize_heading(text,level):
    """标题编号规范化: 一级 "{n}."、其余 "{n}.{n}"(无尾点)。
    model 中的标题已自带编号(如 "1 总论"), 这里只统一分隔格式。"""
    s=str(text or '').strip()
    m=re.match(r'^(\d+(?:\.\d+)*)[\s.、]*(.+)$',s)
    if not m: return s
    num,rest=m.group(1),m.group(2).strip()
    prefix=f'{num}.' if level==1 and '.' not in num else num
    return f'{prefix} {rest}'

class TableNumbering:
    """表编号重排: 配置 pattern "表{chapter}.{index}" -> 按章累计。
    原 caption 形如 "表 3.1-1 xxx"(节-序号), 重排为 "表3.1 xxx"(章.序号)。"""
    def __init__(self): self._count={}
    def caption(self,caption):
        s=str(caption or '').strip()
        m=re.match(r'^表\s*(\d+(?:\.\d+)?)\s*-\s*\d+\s*(.*)$',s)
        if not m: return s
        chapter=m.group(1).split('.')[0]
        self._count[chapter]=self._count.get(chapter,0)+1
        return f'表{chapter}.{self._count[chapter]} {m.group(2)}'.strip()

def _field_run(p,instr,size=9):
    r=p.add_run(); set_font(r,size)
    b=OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'),'begin')
    i=OxmlElement('w:instrText'); i.set(qn('xml:space'),'preserve'); i.text=instr
    e=OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'),'end')
    r._r.append(b); r._r.append(i); r._r.append(e)

def _auto_update_fields(doc):
    """toc.auto_update: 打开文档时提示更新域。"""
    s=doc.settings.element
    if s.find(qn('w:updateFields')) is None:
        el=OxmlElement('w:updateFields'); el.set(qn('w:val'),'true'); s.append(el)

def _render_cover(doc,cover):
    """cover.layout=center: company_name / project_name / report_title / project_code。
    字号为推测项: 公司名16 / 项目名22粗 / 报告标题26粗 / 项目编号12。"""
    company=str(cover.get('company_name') or '').strip()
    pname=str(cover.get('project_name') or cover.get('title') or '').strip()
    rtitle=str(cover.get('report_title') or cover.get('subtitle') or '可行性研究报告').strip()
    code=str(cover.get('project_code') or '').strip()
    def cline(text,size,bold):
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after=Pt(0)
        r=p.add_run(text); set_font(r,size,bold)
    for _ in range(2): doc.add_paragraph()
    if company and company!='待明确': cline(company,16,False)
    for _ in range(5): doc.add_paragraph()
    if pname: cline(pname,22,True)
    cline(rtitle,26,True)
    for _ in range(5): doc.add_paragraph()
    if code: cline(f'项目编号：{code}',12,False)

def _setup_body_section(doc,project_name):
    """封面节无页眉页脚; 正文节页眉 "{project_name} 可行性研究报告"、页脚页码居中、页码从1起。"""
    sec=doc.add_section(WD_SECTION.NEW_PAGE)
    pg=OxmlElement('w:pgNumType'); pg.set(qn('w:start'),'1'); sec._sectPr.append(pg)
    sec.header.is_linked_to_previous=False
    hp=sec.header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=hp.add_run(f'{project_name} 可行性研究报告'); set_font(r,9)
    sec.footer.is_linked_to_previous=False
    fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    _field_run(fp,'PAGE',9)
    return sec

def _add_toc(doc,entries):
    """TOC: 预渲染 1-3 级条目，每个条目含页码（顺序编号）和点线前导符。
    一级条目（大章节）加粗。内嵌 TOC 域，打开文档时 Word 自动更新为真实页码。"""
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run('目  录'); set_font(r,16,True)
    if not entries:
        ph=doc.add_paragraph(); ph.alignment=WD_ALIGN_PARAGRAPH.CENTER
        rr=ph.add_run('（目录将在打开文档时自动更新）'); set_font(rr,BODY_SIZE)
        return
    def fld_char(type_):
        el=OxmlElement('w:fldChar'); el.set(qn('w:fldCharType'),type_); return el
    def instr_run(p_,instr):
        r_=p_.add_run(); set_font(r_,BODY_SIZE)
        i=OxmlElement('w:instrText'); i.set(qn('xml:space'),'preserve'); i.text=instr
        r_._r.append(i)
    for idx,(level,text) in enumerate(entries):
        ep=doc.add_paragraph()
        ep.paragraph_format.left_indent=Mm(7.4*(level-1))
        ep.paragraph_format.line_spacing=1.5; ep.paragraph_format.space_after=Pt(0)
        # 右对齐页码制表位（点线前导符）：所有层级页码统一对齐到正文右边界（可用宽度 146.6mm），
        # 不随 left_indent 层级缩进左移，否则 2/3 级条目的页码会逐级偏离右边界
        ep.paragraph_format.tab_stops.add_tab_stop(Mm(146.6),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
        is_l1=(level==1)
        if idx==0:
            r_=ep.add_run(); set_font(r_,BODY_SIZE); r_._r.append(fld_char('begin'))
            instr_run(ep,r'TOC \o "1-3" \h \z \u')
            r_=ep.add_run(); set_font(r_,BODY_SIZE); r_._r.append(fld_char('separate'))
        r_=ep.add_run(text); set_font(r_,BODY_SIZE,is_l1)
        # 制表符+页码
        r_=ep.add_run(f'\t{idx+1}'); set_font(r_,BODY_SIZE)
        if idx==len(entries)-1:
            r_=ep.add_run(); set_font(r_,BODY_SIZE); r_._r.append(fld_char('end'))

def add_para(doc,text,indent=True,bold=False):
    p=doc.add_paragraph(); p.paragraph_format.line_spacing=1.5; p.paragraph_format.space_after=Pt(0)
    p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent: p.paragraph_format.first_line_indent=Mm(7.4)
    r=p.add_run(str(text)); set_font(r,BODY_SIZE,bold); return p

def add_heading(doc,text,level):
    h=doc.add_heading(normalize_heading(text,level),level)
    if level==1: h.paragraph_format.page_break_before=True
    return h

def add_table(doc,block,numbering=None):
    cols=block.get('columns') or []; rows=block.get('rows') or []
    caption=block.get('caption'); table_id=block.get('table_id')
    if caption:
        text=numbering.caption(caption) if numbering else caption
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(2)
        r=p.add_run(text); set_font(r,BODY_SIZE,True)
    fs=_table_font_size(len(cols))
    t=doc.add_table(rows=1,cols=len(cols)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'; t.autofit=False
    widths=widths_for(len(cols),table_id)
    for i,c in enumerate(cols): set_cell_text(t.rows[0].cells[i],c,True,fs); cell_width(t.rows[0].cells[i],widths[i])
    repeat_header(t.rows[0])
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row): set_cell_text(cells[i],v,False,fs); cell_width(cells[i],widths[i] if i<len(widths) else widths[-1])
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

def _cover_fields(cover):
    company=str(cover.get('company_name') or '').strip()
    if company=='待明确': company=''
    pname=str(cover.get('project_name') or cover.get('title') or '').strip()
    rtitle=str(cover.get('report_title') or cover.get('subtitle') or '可行性研究报告').strip()
    code=str(cover.get('project_code') or '').strip()
    return company,pname,rtitle,code

def export_model(model,outpath):
    doc=configure(Document())
    cover=model.get('cover') or {}
    _render_cover(doc,cover)
    company,pname,rtitle,code=_cover_fields(cover)
    _setup_body_section(doc,pname or '改造项目')
    _add_toc(doc,_collect_toc(model))
    numbering=TableNumbering()
    for sec in model.get('sections') or []:
        heading=sec.get('heading'); level=int(sec.get('level',1))
        if heading: add_heading(doc,heading,level)
        for b in sec.get('blocks') or []:
            typ=b.get('type')
            if typ=='paragraph': add_para(doc,b.get('text',''),b.get('indent',True),b.get('bold',False))
            elif typ=='numbered_list':
                for i,x in enumerate(b.get('items') or [],1): add_para(doc,f'{i}）{x}',False)
            elif typ=='table': add_table(doc,b,numbering)
            elif typ=='heading': add_heading(doc,b.get('text',''),int(b.get('level',2)))
            elif typ=='page_break': doc.add_page_break()
    _auto_update_fields(doc)
    props=model.get('properties') or {}; doc.core_properties.title=props.get('title','可行性研究报告'); doc.core_properties.subject=props.get('subject','流程工业改造项目可行性研究报告'); doc.core_properties.author=''
    doc.save(outpath)
    return outpath


# ===== Markdown 导出（与 DOCX 共用标题编号/表编号规则，正文一致） =====


def _text(value):
    return '' if value is None else str(value)

def _escape_table_cell(value):
    return _text(value).replace('\\', '\\\\').replace('|', '\\|').replace('\r\n', '<br>').replace('\n', '<br>')

def _heading(text, level):
    level=max(1,min(int(level or 1),6))
    return f"{'#'*level} {normalize_heading(text,level)}".rstrip()

def _collect_toc(model):
    entries=[]
    for sec in model.get('sections') or []:
        level=int(sec.get('level',1))
        if sec.get('heading') and level<=3:
            entries.append((level,normalize_heading(sec['heading'],level)))
        for b in sec.get('blocks') or []:
            if b.get('type')=='heading':
                level=int(b.get('level',2))
                if b.get('text') and level<=3:
                    entries.append((level,normalize_heading(b['text'],level)))
    return entries

def render_model(model):
    lines=[]
    cover=model.get('cover') or {}
    company,pname,rtitle,code=_cover_fields(cover)
    if company: lines.append(f'**{company}**'); lines.append('')
    if pname: lines.append(f'# {pname}'); lines.append('')
    lines.append(f'**{rtitle}**'); lines.append('')
    if code: lines.append(f'项目编号：{code}'); lines.append('')
    lines.extend(['---',''])
    toc=_collect_toc(model)
    if toc:
        lines.append('## 目录'); lines.append('')
        for level,title in toc:
            lines.append(f"{'  '*(level-1)}- {title}")
        lines.extend(['','---',''])
    numbering=TableNumbering()
    for sec in model.get('sections') or []:
        heading=sec.get('heading')
        if heading:
            lines.append(_heading(heading,sec.get('level',1)))
            lines.append('')
        for block in sec.get('blocks') or []:
            typ=block.get('type')
            if typ=='paragraph':
                text=_text(block.get('text','')).strip()
                if text:
                    lines.append(text)
                    lines.append('')
            elif typ=='numbered_list':
                for i,item in enumerate(block.get('items') or [],1):
                    lines.append(f'{i}. {_text(item)}')
                if block.get('items'):
                    lines.append('')
            elif typ=='table':
                caption=_text(block.get('caption','')).strip()
                if caption:
                    lines.append(f'**{numbering.caption(caption)}**')
                    lines.append('')
                cols=block.get('columns') or []
                rows=block.get('rows') or []
                if cols:
                    lines.append('| ' + ' | '.join(_escape_table_cell(c) for c in cols) + ' |')
                    lines.append('| ' + ' | '.join('---' for _ in cols) + ' |')
                    for row in rows:
                        vals=list(row or [])
                        if len(vals)<len(cols): vals += ['']*(len(cols)-len(vals))
                        vals=vals[:len(cols)]
                        lines.append('| ' + ' | '.join(_escape_table_cell(v) for v in vals) + ' |')
                    lines.append('')
            elif typ=='heading':
                text=_text(block.get('text','')).strip()
                if text:
                    lines.append(_heading(text,block.get('level',2)))
                    lines.append('')
            elif typ=='page_break':
                # Markdown has no page semantics; keep a visible structural separator.
                lines.extend(['---',''])
    while lines and lines[-1]=='':
        lines.pop()
    return '\n'.join(lines) + '\n'


def export_model_markdown(model,outpath):
    out=Path(outpath)
    out.parent.mkdir(parents=True,exist_ok=True)
    out.write_text(render_model(model),encoding='utf-8')
    return str(out)
